"""Windows acceptance driver for SMOKE-03 (offline package) and RULE-05 (S-规模 within 30 minutes at 8 GB).

Runs on a Windows x64 machine (a CI Windows runner in practice) with its own Python used only as the test driver:

    python packaging/win_acceptance.py --zip dist/机会判定台-1.0.0-win64.zip --fixtures fixtures/packages --out results

What it does
- SMOKE-03: unzip the package; block outbound network for the package's python.exe with a firewall rule; run every
  package command with a clean environment whose PATH has no Python; double-click equivalent of 启动.bat; check that
  the service answers on 127.0.0.1 only, that the default browser was asked to open the workbench, that `oppo --help`
  works inside the install directory; import H1-A; stop; start again and check the imported package is still there and
  that data lives in data/ apart from the program; a second start while running only opens the browser.
- RULE-05: run import → contract confirm → compute → backtest on S-规模 with the package's python.exe inside a Windows
  Job Object limited to 8 GB of committed memory; record start/end times, durations and peak memory.
Results go to <out>/results.json, logs and a screenshot of the workbench.
"""

from __future__ import annotations

import argparse
import ctypes
import json
import os
import shutil
import socket
import subprocess
import sys
import time
import urllib.request
import zipfile
from ctypes import wintypes
from datetime import datetime
from pathlib import Path

PORT = 8765


def clean_env(root: Path) -> dict[str, str]:
    sysroot = os.environ.get("SystemRoot", r"C:\Windows")
    env = {k: os.environ[k] for k in ("SystemRoot", "SystemDrive", "TEMP", "TMP", "USERPROFILE", "LOCALAPPDATA",
                                      "APPDATA", "COMPUTERNAME", "USERNAME", "ComSpec", "PATHEXT", "windir",
                                      "ProgramData", "PROGRAMFILES", "HOMEDRIVE", "HOMEPATH") if k in os.environ}
    env["PATH"] = rf"{sysroot}\System32;{sysroot};{sysroot}\System32\Wbem;{sysroot}\System32\WindowsPowerShell\v1.0"
    return env


def run(cmd: list[str] | str, env: dict[str, str], cwd: Path, log: Path, timeout: int = 3600) -> tuple[int, str]:
    p = subprocess.run(cmd, env=env, cwd=str(cwd), capture_output=True, timeout=timeout, shell=isinstance(cmd, str))
    out = (p.stdout or b"").decode("utf-8", "replace") + (p.stderr or b"").decode("utf-8", "replace")
    with log.open("a", encoding="utf-8") as f:
        f.write(f"\n$ {cmd}\n[exit {p.returncode}]\n{out}\n")
    return p.returncode, out


def wait_http(url: str, seconds: int = 120) -> bool:
    end = time.time() + seconds
    while time.time() < end:
        try:
            with urllib.request.urlopen(url, timeout=3) as r:
                if r.status == 200:
                    return True
        except OSError:
            time.sleep(1)
    return False


def listening_addresses(port: int) -> list[str]:
    out = subprocess.run(["netstat", "-ano", "-p", "TCP"], capture_output=True, text=True).stdout
    out6 = subprocess.run(["netstat", "-ano", "-p", "TCPv6"], capture_output=True, text=True).stdout
    addrs = []
    for line in (out + out6).splitlines():
        parts = line.split()
        if len(parts) >= 4 and parts[3].upper() == "LISTENING" and parts[1].endswith(f":{port}"):
            addrs.append(parts[1].rsplit(":", 1)[0])
    return addrs


def lan_ip() -> str | None:
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("10.255.255.255", 1))
        ip = s.getsockname()[0]
    except OSError:
        ip = None
    finally:
        s.close()
    return ip if ip and not ip.startswith("127.") else None


def browser_requests(port: int) -> list[str]:
    ps = ("Get-CimInstance Win32_Process | Where-Object { $_.CommandLine -like '*127.0.0.1:%d*' } | "
          "ForEach-Object { $_.Name + ' :: ' + $_.CommandLine }" % port)
    out = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True).stdout
    return [ln for ln in out.splitlines() if ln.strip() and "python" not in ln.lower()]


def kill_tree(pid: int) -> None:
    subprocess.run(["taskkill", "/T", "/F", "/PID", str(pid)], capture_output=True)


def screenshot(url: str, path: Path) -> list[str]:
    from playwright.sync_api import sync_playwright

    errors: list[str] = []
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        pg = b.new_page(viewport={"width": 1440, "height": 900})
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
        pg.goto(url, wait_until="networkidle")
        pg.wait_for_timeout(800)
        pg.screenshot(path=str(path), full_page=True)
        b.close()
    return errors


# ---- Job Object with a memory limit (RULE-05) -----------------------------------------------------

class IO_COUNTERS(ctypes.Structure):
    _fields_ = [(n, ctypes.c_ulonglong) for n in ("ReadOperationCount", "WriteOperationCount", "OtherOperationCount",
                                                    "ReadTransferCount", "WriteTransferCount", "OtherTransferCount")]


class JOBOBJECT_BASIC_LIMIT_INFORMATION(ctypes.Structure):
    _fields_ = [("PerProcessUserTimeLimit", ctypes.c_longlong), ("PerJobUserTimeLimit", ctypes.c_longlong),
                ("LimitFlags", wintypes.DWORD), ("MinimumWorkingSetSize", ctypes.c_size_t),
                ("MaximumWorkingSetSize", ctypes.c_size_t), ("ActiveProcessLimit", wintypes.DWORD),
                ("Affinity", ctypes.c_size_t), ("PriorityClass", wintypes.DWORD), ("SchedulingClass", wintypes.DWORD)]


class JOBOBJECT_EXTENDED_LIMIT_INFORMATION(ctypes.Structure):
    _fields_ = [("BasicLimitInformation", JOBOBJECT_BASIC_LIMIT_INFORMATION), ("IoInfo", IO_COUNTERS),
                ("ProcessMemoryLimit", ctypes.c_size_t), ("JobMemoryLimit", ctypes.c_size_t),
                ("PeakProcessMemoryUsed", ctypes.c_size_t), ("PeakJobMemoryUsed", ctypes.c_size_t)]


JOB_OBJECT_LIMIT_JOB_MEMORY = 0x200
JobObjectExtendedLimitInformation = 9
CREATE_SUSPENDED = 0x4


class MemoryJob:
    def __init__(self, limit_bytes: int):
        k32 = ctypes.WinDLL("kernel32", use_last_error=True)
        k32.CreateJobObjectW.restype = wintypes.HANDLE
        self.k32 = k32
        self.ntdll = ctypes.WinDLL("ntdll")
        self.job = k32.CreateJobObjectW(None, None)
        info = JOBOBJECT_EXTENDED_LIMIT_INFORMATION()
        info.BasicLimitInformation.LimitFlags = JOB_OBJECT_LIMIT_JOB_MEMORY
        info.JobMemoryLimit = limit_bytes
        ok = k32.SetInformationJobObject(wintypes.HANDLE(self.job), JobObjectExtendedLimitInformation,
                                         ctypes.byref(info), ctypes.sizeof(info))
        if not ok:
            raise OSError(ctypes.get_last_error(), "SetInformationJobObject failed")

    def run(self, cmd: list[str], env: dict[str, str], cwd: Path, log: Path) -> tuple[int, str, float]:
        t0 = time.time()
        p = subprocess.Popen(cmd, env=env, cwd=str(cwd), stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                             creationflags=CREATE_SUSPENDED)
        if not self.k32.AssignProcessToJobObject(wintypes.HANDLE(self.job), wintypes.HANDLE(int(p._handle))):
            raise OSError(ctypes.get_last_error(), "AssignProcessToJobObject failed")
        self.ntdll.NtResumeProcess(wintypes.HANDLE(int(p._handle)))
        out, _ = p.communicate()
        text = out.decode("utf-8", "replace")
        with log.open("a", encoding="utf-8") as f:
            f.write(f"\n$ {cmd}\n[exit {p.returncode}] {time.time() - t0:.1f}s\n{text}\n")
        return p.returncode, text, time.time() - t0

    def peaks(self) -> dict[str, float]:
        info = JOBOBJECT_EXTENDED_LIMIT_INFORMATION()
        self.k32.QueryInformationJobObject(wintypes.HANDLE(self.job), JobObjectExtendedLimitInformation,
                                           ctypes.byref(info), ctypes.sizeof(info), None)
        return {"peak_job_memory_mb": round(info.PeakJobMemoryUsed / 2**20, 1),
                "peak_process_memory_mb": round(info.PeakProcessMemoryUsed / 2**20, 1),
                "job_memory_limit_mb": round(info.JobMemoryLimit / 2**20, 1)}


# ---- main ---------------------------------------------------------------------------------------------


def smoke03(zip_path: Path, fixtures: Path, out: Path, res: dict) -> Path:
    log = out / "smoke03.log"
    base = Path(os.environ.get("SystemDrive", "C:") + "\\") / "oppo-acceptance"
    if base.exists():
        shutil.rmtree(base, ignore_errors=True)
    base.mkdir(parents=True)
    with zipfile.ZipFile(zip_path) as zf:
        zf.extractall(base)
    root = next(p for p in base.iterdir() if p.is_dir())
    py = root / "python" / "python.exe"
    s: dict = {"install_dir": str(root), "zip": zip_path.name, "zip_mb": round(zip_path.stat().st_size / 1e6, 1)}
    res["SMOKE-03"] = s
    fw = subprocess.run(["netsh", "advfirewall", "firewall", "add", "rule", "name=oppo-offline-block", "dir=out",
                         "action=block", f"program={py}", "enable=yes"], capture_output=True, text=True)
    s["firewall_block_outbound"] = fw.returncode == 0
    env = clean_env(root)
    code, where = run("where python", env, root, log)
    s["python_on_clean_path"] = code == 0
    s["python_on_clean_path_output"] = where.strip()
    code, help_out = run(["cmd", "/c", str(root / "oppo.bat"), "--help"], env, root, log)
    s["oppo_help_exit"] = code
    s["oppo_help_commands"] = [c for c in ("import", "contract", "health", "merge", "compute", "backtest", "ranking",
                                           "decision", "export", "restore") if c in help_out]
    t0 = time.time()
    log1 = (out / "server_first.log").open("wb")
    srv = subprocess.Popen(["cmd", "/c", str(root / "启动.bat")], env=env, cwd=str(root),
                           stdout=log1, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                           creationflags=subprocess.CREATE_NEW_PROCESS_GROUP)
    s["first_start_ok"] = wait_http(f"http://127.0.0.1:{PORT}/")
    s["first_start_seconds"] = round(time.time() - t0, 1)
    s["first_start_process_exit"] = srv.poll()
    log1.flush()
    s["first_start_output"] = (out / "server_first.log").read_bytes().decode("utf-8", "replace")[-1500:]
    if not s["first_start_ok"]:
        kill_tree(srv.pid)
        code, diag = run([str(root / "python" / "python.exe"), "-c", "import oppo.web.server, uvicorn; print('import ok')"],
                         env, root, log, timeout=120)
        s["diagnose_import"] = diag.strip()[-1500:]
        return root
    time.sleep(3)
    s["listening"] = listening_addresses(PORT)
    ip = lan_ip()
    if ip:
        try:
            urllib.request.urlopen(f"http://{ip}:{PORT}/", timeout=3)
            s["lan_access"] = "reachable"
        except OSError as exc:
            s["lan_access"] = f"refused ({exc.__class__.__name__})"
    s["browser_launch_processes"] = browser_requests(PORT)
    s["workbench_console_errors"] = screenshot(f"http://127.0.0.1:{PORT}/", out / "smoke03_workbench.png")
    code, imp = run(["cmd", "/c", str(root / "oppo.bat"), "import", str(fixtures / "H1-A")], env, root, log)
    s["import_exit"] = code
    s["import_output"] = imp.strip()[-300:]
    code, again = run(["cmd", "/c", str(root / "启动.bat"), "--no-browser"], env, root, log, timeout=60)
    s["second_start_while_running"] = again.strip()[-200:]
    kill_tree(srv.pid)
    time.sleep(3)
    log2 = (out / "server_second.log").open("wb")
    srv2 = subprocess.Popen(["cmd", "/c", str(root / "启动.bat")], env=env, cwd=str(root),
                            stdout=log2, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                            creationflags=subprocess.CREATE_NEW_PROCESS_GROUP)
    s["restart_ok"] = wait_http(f"http://127.0.0.1:{PORT}/")
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{PORT}/packages", timeout=10) as r:
            page = r.read().decode("utf-8")
        s["package_listed_after_restart"] = "P-0001" in page
    except OSError as exc:
        s["package_listed_after_restart"] = f"error {exc}"
    s["restart_console_errors"] = screenshot(f"http://127.0.0.1:{PORT}/packages", out / "smoke03_packages_after_restart.png")
    kill_tree(srv2.pid)
    s["data_dir_files"] = sorted(p.name for p in (root / "data").iterdir())
    s["program_dirs"] = sorted(p.name for p in root.iterdir() if p.is_dir())
    s["data_separate_from_program"] = (root / "data" / "oppo.sqlite3").exists() and not (root / "app" / "oppo.sqlite3").exists()
    subprocess.run(["netsh", "advfirewall", "firewall", "delete", "rule", "name=oppo-offline-block"], capture_output=True)
    res["SMOKE-03"] = s
    return root


def rule05(root: Path, scale: Path, out: Path, res: dict) -> None:
    log = out / "rule05.log"
    env = clean_env(root)
    data = root.parent / "rule05-data"
    if data.exists():
        shutil.rmtree(data, ignore_errors=True)
    env.update(OPPO_DATA_DIR=str(data), PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
    py = str(root / "python" / "python.exe")
    job = MemoryJob(8 * 2**30)
    r: dict = {"machine": {"cpu_count": os.cpu_count(), "platform": sys.platform,
                           "total_memory_gb": None}, "steps": {}}
    try:
        mem = subprocess.run(["powershell", "-NoProfile", "-Command",
                              "(Get-CimInstance Win32_ComputerSystem).TotalPhysicalMemory"], capture_output=True, text=True)
        r["machine"]["total_memory_gb"] = round(int(mem.stdout.strip()) / 2**30, 1)
        osv = subprocess.run(["cmd", "/c", "ver"], capture_output=True, text=True)
        r["machine"]["windows"] = osv.stdout.strip()
    except (ValueError, OSError):
        pass
    steps = [("import", ["import", str(scale)]), ("contract_confirm", ["contract", "confirm"]),
             ("compute", ["compute"]), ("backtest", ["backtest"])]
    for name, args in steps:
        started = datetime.now().isoformat(timespec="seconds")
        code, text, secs = job.run([py, "-m", "oppo.cli", *args], env, root, log)
        r["steps"][name] = {"exit": code, "started": started, "ended": datetime.now().isoformat(timespec="seconds"),
                            "seconds": round(secs, 1), "memory_error": "MemoryError" in text or "内存" in text,
                            "tail": text.strip()[-300:]}
        if code != 0:
            break
    r.update(job.peaks())
    st = r["steps"]
    if "compute" in st and "backtest" in st:
        r["compute_to_backtest_seconds"] = round(st["compute"]["seconds"] + st["backtest"]["seconds"], 1)
        r["within_30_minutes"] = r["compute_to_backtest_seconds"] <= 1800 and all(x["exit"] == 0 for x in st.values())
    res["RULE-05"] = r


def flow17(root: Path, out: Path, res: dict) -> None:
    """FLOW-17 (browser part): finish the month on the package's own CLI, copy the exported web report to another
    folder and open it with the network off in Chromium and in WebKit (Safari's engine)."""
    log = out / "flow17.log"
    env = clean_env(root)
    r: dict = {}
    res["FLOW-17"] = r
    oppo = ["cmd", "/c", str(root / "oppo.bat")]
    for args in (["contract", "confirm"], ["merge", "confirm", "--basis", "标题主干相同"], ["compute"], ["backtest"],
                 ["backup-dir", str(root.parent / "backups")]):
        code, _ = run(oppo + args, env, root, log)
        r["cli_" + args[0]] = code
    code, text = run(oppo + ["export"], env, root, log)
    r["cli_export"] = code
    info = json.loads(text[text.index("{"):text.rindex("}") + 1])
    copy_dir = root.parent / "report-copy"
    copy_dir.mkdir(exist_ok=True)
    html = copy_dir / Path(info["html_path"]).name
    shutil.copy2(info["html_path"], html)
    r["report_copy"] = str(html)
    r["backup"] = (info.get("backup") or {}).get("status")
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright

    wb = load_workbook(info["xlsx_path"], read_only=True)
    r["sheets"] = wb.sheetnames
    ws = wb["细分排序"]
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, row in enumerate(rows) if row and row[0] == "排名")
    excel_order = ["".join(str(x) for x in row[2:5]) for row in rows[head + 1:] if row and row[1]]
    norm = lambda t: t.replace(" ", "").replace("·", "").replace("\n", "")  # noqa: E731
    texts = {}
    with sync_playwright() as pw:
        for name in ("chromium", "webkit"):
            b = getattr(pw, name).launch()
            ctx = b.new_context(viewport={"width": 1440, "height": 900}, offline=True)
            pg = ctx.new_page()
            reqs: list[str] = []
            pg.on("request", lambda q, reqs=reqs: reqs.append(q.url))
            pg.goto(html.as_uri())
            pg.wait_for_timeout(800)
            pg.evaluate("document.querySelectorAll('details').forEach(d => d.open = true)")
            order = pg.evaluate("""() => { const t = [...document.querySelectorAll('table')].find(
                x => x.querySelector('th') && [...x.querySelectorAll('th')].some(h => h.textContent.trim() === '得分'));
                return t ? [...t.querySelectorAll('tbody tr')].filter(tr => !tr.classList.contains('table__divider'))
                  .map(tr => tr.children[1].innerText) : []; }""")
            texts[name] = pg.evaluate("document.body.innerText")
            pg.screenshot(path=str(out / f"flow17_report_{name}.png"), full_page=True)
            r[name] = {"network_requests": [u for u in reqs if not u.startswith("file:")],
                       "ranking_rows": len(order),
                       "ranking_matches_excel": [norm(x) for x in order] == [norm(x) for x in excel_order]}
            ctx.close()
            b.close()
    r["same_text_in_both_engines"] = texts.get("chromium") == texts.get("webkit")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--zip", required=True)
    ap.add_argument("--fixtures", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    out = Path(a.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    res: dict = {"started": datetime.now().isoformat(timespec="seconds")}
    root = None
    try:
        try:
            root = smoke03(Path(a.zip).resolve(), Path(a.fixtures).resolve(), out, res)
        except Exception as exc:  # noqa: BLE001 - recorded, RULE-05 still runs
            res.setdefault("SMOKE-03", {})["error"] = repr(exc)
            root = next((p for p in (Path(os.environ.get("SystemDrive", "C:") + "\\") / "oppo-acceptance").iterdir()
                         if p.is_dir()), None)
        if root is not None:
            try:
                flow17(root, out, res)
            except Exception as exc:  # noqa: BLE001 - recorded; RULE-05 still runs
                res.setdefault("FLOW-17", {})["error"] = repr(exc)
            rule05(root, Path(a.fixtures).resolve() / "S-规模", out, res)
    finally:
        res["ended"] = datetime.now().isoformat(timespec="seconds")
        (out / "results.json").write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
