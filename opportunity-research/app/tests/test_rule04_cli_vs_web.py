"""RULE-04: database A done only with the command line, database B with the same inputs through the web pages;
results must be identical and every business record must carry operator and time."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from oppo.core import backup as BK
from tests.conftest import APP_ROOT, fixture_pkg
from tests.test_e2e_smoke02 import month_in_browser

RESULT_QUERIES = {
    "spu": "SELECT * FROM spu WHERE run_id=(SELECT MAX(id) FROM compute_run WHERE mode='history') ORDER BY spu_key",
    "judgment": "SELECT * FROM judgment WHERE run_id=(SELECT MAX(id) FROM compute_run WHERE mode='history') "
                "ORDER BY spu_key",
    "signal_value": "SELECT * FROM signal_value WHERE run_id=(SELECT MAX(id) FROM compute_run WHERE mode='history') "
                    "ORDER BY t_month, segment_key, signal",
    "backtest": "SELECT code, mode, conclusion, reason_class, seed, result_json FROM backtest_run ORDER BY id",
    "card": "SELECT label, mode, status, signals_json, basis_json FROM decision_card ORDER BY id",
    "rules": "SELECT r.code, r.origin, v.version_no, v.condition_json, v.action, v.status, v.basis_json "
             "FROM rule r JOIN rule_version v ON v.rule_id = r.id ORDER BY r.code, v.version_no",
    "snapshot": "SELECT month, seq, mode, card_label, card_status, sha256, content_json FROM snapshot ORDER BY id",
    "merge": "SELECT parents_json, status FROM merge ORDER BY parents_json",
    "contract": "SELECT file_code, field, meaning, grain, timing, nature, usage, status FROM contract_item "
                "ORDER BY file_code, field",
}
AUDITED = ("package", "contract_version", "merge", "compute_run", "backtest_run", "decision_card", "rule_version",
           "snapshot", "export_record", "backup_record", "run_step", "change_log")
DROP = {"run_id", "id"}


def _rows(db: Path) -> dict[str, list[tuple]]:
    import sqlite3

    c = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    out = {}
    for name, q in RESULT_QUERIES.items():
        out[name] = [tuple(v for k, v in zip(r.keys(), tuple(r), strict=True) if k not in DROP) for r in c.execute(q)]
    c.close()
    return out


def _audit_gaps(db: Path) -> list[str]:
    import sqlite3

    c = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    gaps = []
    for t in AUDITED:
        n = c.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        bad = c.execute(f"SELECT COUNT(*) FROM {t} WHERE created_by IS NULL OR created_by = '' "
                        f"OR created_at IS NULL OR created_at = ''").fetchone()[0]
        if n == 0 or bad:
            gaps.append(f"{t}: rows={n} missing={bad}")
    c.close()
    return gaps


def _sheets(path: str, root: Path) -> dict[str, list[tuple]]:
    """Cell values with times and this database's own paths masked (they differ by nature)."""
    import re

    stamp = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?")

    def norm(v):  # type: ignore[no-untyped-def]
        if isinstance(v, str):
            v = stamp.sub("<time>", v.replace(str(root), "<root>"))
            return re.sub(r"<time> · [\d.]+ ?(秒|分钟|小时)", "<time> · <duration>", v)
        return v

    wb = load_workbook(path, read_only=True)
    return {ws.title: [tuple(norm(v) for v in r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def test_rule04_cli_and_web_give_the_same_results(live_server, tmp_path, page) -> None:  # type: ignore[no-untyped-def]
    a = tmp_path / "A"
    env_py = [sys.executable, "-m", "oppo.cli", "--data-dir", str(a)]
    steps = [["import", str(fixture_pkg("H1-A"))], ["contract", "confirm", "--suggested"],
             ["merge", "confirm", "--basis", "标题主干相同"], ["compute"], ["backtest"],
             ["backup-dir", str(tmp_path / "备份A")], ["export"]]
    out = ""
    for st in steps:
        p = subprocess.run(env_py + st, cwd=str(APP_ROOT), capture_output=True, text=True)
        assert p.returncode == 0, (st, p.stdout, p.stderr)
        out = p.stdout
    xlsx_a = json.loads(out[out.index("{"):])["xlsx_path"]

    b = tmp_path / "B"
    b.mkdir()
    live = live_server(b)
    with live.app.engine.begin() as conn:
        BK.set_directory(conn, str(tmp_path / "备份B"), "wen")
    month_in_browser(live, page)
    with live.app.engine.begin() as conn:
        xlsx_b = conn.execute(text("SELECT xlsx_path FROM export_record ORDER BY id DESC")).scalar()

    ra, rb = _rows(a / "oppo.sqlite3"), _rows(b / "oppo.sqlite3")
    for name in RESULT_QUERIES:
        assert ra[name] == rb[name], name
    assert _audit_gaps(a / "oppo.sqlite3") == [] and _audit_gaps(b / "oppo.sqlite3") == []
    sa, sb = _sheets(xlsx_a, tmp_path / "A"), _sheets(xlsx_b, tmp_path / "B")
    assert list(sa) == list(sb) and len(sa) == 9
    for title in sa:
        diff = [(i, x, y) for i, (x, y) in enumerate(zip(sa[title], sb[title], strict=False)) if x != y]
        assert len(sa[title]) == len(sb[title]) and not diff, (title, diff[:3])
