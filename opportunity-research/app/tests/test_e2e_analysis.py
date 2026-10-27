"""Browser E2E for the analysis pages: 回测 (PAGE-F07-01), 判定卡 (F08-01), 立项锦囊 (F08-02), 本月排序 (F09-01),
导出报告弹窗 (F12-01) and the offline web report (F12-02).

Expected numbers come from the H1-A / X1 / H1-无区分度 sample-package documents (fixtures/packages/*.说明.md),
not from the implementation under test."""

from __future__ import annotations

import re
import sqlite3
from pathlib import Path

import pytest
from sqlalchemy import select

from oppo import db as D
from oppo.core import backup as BK
from oppo.core import cards as CARD
from tests.conftest import _prepared

X_SEG = "30–48″ · 落地式 · $500–900"
Y_SEG = "30–48″ · 落地式 · <$500"


@pytest.fixture(scope="module")
def x1_done(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return _prepared(tmp_path_factory, "x1_done", ("import:X1", "confirm", "compute", "backtest:cross"))


@pytest.fixture(scope="module")
def nd_done(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """H1-无区分度：回测未达标（信号无区分度），判定卡未验证。"""
    return _prepared(tmp_path_factory, "nd_done", ("import:H1-无区分度", "confirm", "compute", "backtest"))


def _no_errors(page) -> None:  # type: ignore[no-untyped-def]
    assert not page.js_errors, page.js_errors


def test_backtest_page_history(live_server, clone, h1a_done, page):  # type: ignore[no-untyped-def]
    live = live_server(clone(h1a_done))
    page.goto(live.u("/backtest"))
    concl = page.locator('[data-concl="history"]')
    assert "BT-0001" in concl.inner_text()
    assert concl.locator(".badge", has_text="达标").count() == 1
    assert concl.locator(".hero-figure").inner_text().strip() == "3.00×"
    txt = concl.inner_text()
    assert "24 / 40" in txt and "60 / 300" in txt and "60.0%" in txt and "20.0%" in txt
    assert "广告投入、站外引流未纳入" in txt
    # split 6 rule batches / 3 test batches / 5 not decided (H1-A: decided 2025-01..2025-09)
    stats = page.locator("#split-title").locator("xpath=../..").locator(".stat__value").all_inner_texts()
    assert [s.split("个")[0].strip() for s in stats] == ["6", "3", "5"]
    card_row = page.locator('tr[data-row="card"]')
    assert "24 / 40" in card_row.inner_text() and "3.00×" in card_row.inner_text()
    base = page.locator('tr[data-row="base"]')
    assert "60 / 300" in base.inner_text() and "倍数基准" in base.inner_text()
    # every current indicator is listed with the 候选 tag
    for label in ("单位 ASIN 贡献", "属性溢价", "属性趋势", "低分高销"):
        assert page.locator("table.table tr", has_text=label).locator(".etag--cand").count() >= 1
    rejected = page.locator("#snap-title").locator("xpath=../..").inner_text()
    for f in ("过去一个月购买", "当月价格", "当月评分"):
        assert f in rejected
    assert "20261106" in page.locator("#run-title").locator("xpath=../..").inner_text()
    _no_errors(page)


def test_card_and_kits(live_server, clone, h1a_done, page):  # type: ignore[no-untyped-def]
    live = live_server(clone(h1a_done))
    page.goto(live.u("/card"))
    banner = page.locator(".banner--ok").first
    assert "已达标 · 判定卡 v1（历史口径）" in banner.inner_text()
    sig = page.locator('tr[data-signal="demand"]').inner_text()
    assert "需求在涨" in sig and "≥ +15%" in sig and "1.00" in sig
    assert page.locator("#card-versions tbody tr").count() == 1
    assert "3.00×（0.82–3.36）" in page.locator("#card-versions tbody tr").inner_text()
    _no_errors(page)

    page.goto(live.u("/kits"))
    tabs = page.locator('[role="tab"]').all_inner_texts()
    assert [re.sub(r"\s+", "", t) for t in tabs] == ["全部3", "立项候选1", "观察1", "放弃1"]
    page.locator('[role="tab"]', has_text="立项候选").click()
    go = page.locator("#kits-go")
    assert go.is_visible() and go.locator("article").count() == 1
    assert "需求在涨 ≥ +15%" in go.inner_text()
    assert go.locator("a.card-link").get_attribute("href") == "/rules?rule=R-001"
    page.locator('[role="tab"]', has_text="放弃").click()
    assert "需求在涨 < 0%" in page.locator("#kits-drop").inner_text()
    _no_errors(page)


def test_ranking_order(live_server, clone, h1a_done, page):  # type: ignore[no-untyped-def]
    live = live_server(clone(h1a_done))
    page.goto(live.u("/ranking"))
    sub = page.locator(".page-head__sub").all_inner_texts()
    assert "2026-03 · 数据至 2026-02" in sub[0] and "判定卡 v1 已达标" in sub[0]
    assert "本月快照已保存（第 1 版）" in sub[1] and "起量新品（任一月父体月销 ≥ 60 件）" in sub[1]
    rows = page.locator("#rank-table tbody tr")
    first, second = rows.nth(0), rows.nth(1)
    # equal score: larger sample first (REQ-F09-01)
    assert first.get_attribute("data-segment") == X_SEG.replace(" · ", "|") and "n=44" in first.inner_text()
    assert second.get_attribute("data-segment") == Y_SEG.replace(" · ", "|") and "n=42" in second.inner_text()
    assert page.locator("#rank-table tbody tr[data-segment]:not([data-insufficient])").count() == 10
    divider = page.locator("#rank-table tr.table__divider")
    assert "样本不足 · 3 个细分" in divider.inner_text()
    # every insufficient row comes after the divider and keeps an interval and its 锦囊
    after = page.locator("#rank-table tr.table__divider ~ tr")
    assert after.count() == 3
    for i in range(3):
        r = after.nth(i)
        assert r.get_attribute("data-insufficient") is not None
        assert r.locator(".ibar--grey").count() == 1 and "样本不足" in r.inner_text()
    ranks = [int(t) for t in page.locator("#rank-table tbody tr[data-segment] td:first-child").all_inner_texts()]
    assert ranks == list(range(1, 14))
    _no_errors(page)


def _export_via_ui(page, live) -> int:  # type: ignore[no-untyped-def]
    page.goto(live.u("/ranking?export=1"))
    modal = page.locator("#export-modal")
    modal.locator("[data-export-start]").first.click()
    page.wait_for_url(re.compile(r"done=\d+"), timeout=120_000)
    return int(re.search(r"done=(\d+)", page.url).group(1))  # type: ignore[union-attr]


def test_export_from_ranking(live_server, clone, h1a_done, page, browser, tmp_path):  # type: ignore[no-untyped-def]
    live = live_server(clone(h1a_done))
    bdir = tmp_path / "backups"
    with live.app.engine.begin() as conn:
        assert BK.set_directory(conn, str(bdir), "wen")[0]
    eid = _export_via_ui(page, live)
    res = page.locator("#exp-result")
    assert res.is_visible()
    txt = res.inner_text()
    assert "导出完成" in txt and "快照已标为“已导出”" in txt and "备份完成" in txt
    with live.app.engine.begin() as conn:
        rec = conn.execute(select(D.export_record).where(D.export_record.c.id == eid)).mappings().first()
        assert rec is not None and rec["backup_id"]
        bk = conn.execute(select(D.backup_record).where(D.backup_record.c.id == rec["backup_id"])).mappings().first()
        assert bk["status"] == "success" and Path(bk["path"]).is_file()
        assert conn.execute(select(D.snapshot_export).where(D.snapshot_export.c.export_id == eid)).first()
        step = conn.execute(select(D.run_step).where(D.run_step.c.step == "导出")).mappings().first()
        assert step is not None and step["month"] == "2026-03"
    xlsx, html = Path(rec["xlsx_path"]), Path(rec["html_path"])
    assert xlsx.name == "浴室柜_机会判定_2026-03_快照1.xlsx" and xlsx.is_file()
    assert html.name == "浴室柜_机会判定_2026-03_快照1.html" and html.is_file()
    assert str(xlsx) in txt and str(html) in txt

    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["首页", "细分排序", "锦囊", "判定卡版本", "回测摘要", "台账-快照", "台账-立项记录",
                             "台账-对账", "版本信息"]
    assert wb["首页"]["A1"].value == "判定卡 v1 已达标（历史口径）"
    rank = wb["细分排序"]
    assert rank["B3"].value == X_SEG and rank["B4"].value == Y_SEG
    for ws in wb.worksheets[1:]:
        assert ws.freeze_panes == "A3"

    doc = html.read_text(encoding="utf-8")
    assert not re.search(r'(src|href)\s*=\s*"(https?:)?//', doc), "report must not reference external resources"
    assert "<script" not in doc and "<link" not in doc and "@import" not in doc and "url(http" not in doc
    assert "判定卡 v1 已达标（历史口径）" in doc and "起量新品" in doc
    # the report opens offline with no request leaving the file
    ctx = browser.new_context(offline=True)
    p2 = ctx.new_page()
    reqs: list[str] = []
    p2.on("request", lambda r: reqs.append(r.url))
    p2.goto(html.as_uri())
    assert p2.locator("h1").inner_text() == "浴室柜 · 2026-03 细分机会排序"
    assert p2.locator("details").count() == 5
    assert [u for u in reqs if not u.startswith("file:")] == []
    ctx.close()

    # second export of the same snapshot never overwrites the first files
    eid2 = _export_via_ui(page, live)
    with live.app.engine.begin() as conn:
        rec2 = conn.execute(select(D.export_record).where(D.export_record.c.id == eid2)).mappings().first()
    assert Path(rec2["xlsx_path"]).name == "浴室柜_机会判定_2026-03_快照1_2.xlsx" and xlsx.is_file()

    page.goto(live.u("/"))
    step8 = page.locator("li.step").nth(7)
    assert step8.get_attribute("data-status") == "done" and "已导出" in step8.inner_text()
    _no_errors(page)


def test_export_without_backup_dir_reports_failure(live_server, clone, h1a_done, page):  # type: ignore[no-untyped-def]
    live = live_server(clone(h1a_done))
    page.goto(live.u("/ranking?export=1"))
    assert "导出前请先设置备份目录" in page.locator("#exp-confirm").inner_text()
    _export_via_ui(page, live)
    item = page.locator('[data-backup="failed"]')
    assert item.is_visible()
    t = item.inner_text()
    assert "备份失败" in t and "未设置备份目录" in t and "上次成功备份" in t and "已有备份未改动" in t
    assert page.locator('[data-file="xlsx"]').is_visible()  # files are still there


def test_unverified_everywhere(live_server, clone, nd_done, page, tmp_path):  # type: ignore[no-untyped-def]
    live = live_server(clone(nd_done))
    with live.app.engine.begin() as conn:
        card = CARD.current_card(conn, live.app.category_id(), "history")
        assert card["status"] == "未验证"
        BK.set_directory(conn, str(tmp_path / "b"), "wen")
    for path in ("/card", "/kits", "/ranking"):
        page.goto(live.u(path))
        assert page.locator(".banner--unverified", has_text=CARD.UNVERIFIED_TEXT).count() >= 1, path
    page.goto(live.u("/backtest"))
    concl = page.locator('[data-concl="history"]').inner_text()
    assert "未达标：信号无区分度" in concl
    page.goto(live.u("/ranking?export=1"))
    assert "判定卡未验证，不可作为立项结论" in page.locator("#exp-confirm").inner_text()
    eid = _export_via_ui(page, live)
    with live.app.engine.begin() as conn:
        rec = conn.execute(select(D.export_record).where(D.export_record.c.id == eid)).mappings().first()
    doc = Path(rec["html_path"]).read_text(encoding="utf-8")
    body = doc[doc.index("<main>"):]
    assert body.index(CARD.UNVERIFIED_TEXT) < body.index("本月结论")
    from openpyxl import load_workbook

    wb = load_workbook(rec["xlsx_path"])
    assert wb["首页"]["A1"].value.startswith(CARD.UNVERIFIED_TEXT)
    for name in ("细分排序", "锦囊", "判定卡版本"):
        assert wb[name]["A1"].value.startswith(CARD.UNVERIFIED_TEXT)
    _no_errors(page)


def test_cross_backtest_and_report(live_server, clone, x1_done, page, tmp_path):  # type: ignore[no-untyped-def]
    live = live_server(clone(x1_done))
    page.goto(live.u("/backtest"))   # only a cross package → cross mode by default
    concl = page.locator('[data-concl="cross"]').inner_text()
    assert "达标（截面）" in concl and CARD.CROSS_TEXT in concl
    assert "3.01×" in concl                       # X1 doc: 10/22 vs 16/106
    assert "10 / 22" in page.locator('tr[data-row="card"]').inner_text()
    assert "16 / 106" in page.locator('tr[data-row="base"]').inner_text()
    heat = page.locator("table.heat")
    assert heat.locator("tbody tr").first.locator(".cell-sub", has_text="前 20%").count() == 1
    assert heat.locator("td.hx", has_text="样本过少").count() >= 1
    page.goto(live.u("/ranking"))
    assert page.locator("#main .banner--cross", has_text=CARD.CROSS_TEXT).count() == 1
    with live.app.engine.begin() as conn:
        BK.set_directory(conn, str(tmp_path / "b"), "wen")
    eid = _export_via_ui(page, live)
    with live.app.engine.begin() as conn:
        rec = conn.execute(select(D.export_record).where(D.export_record.c.id == eid)).mappings().first()
    doc = Path(rec["html_path"]).read_text(encoding="utf-8")
    assert CARD.CROSS_TEXT in doc and CARD.CROSS_LINE_TEXT in doc
    _no_errors(page)


def test_backtest_job_from_page(live_server, clone, h1a_done, page):  # type: ignore[no-untyped-def]
    """运行回测 from the page: background job → new backtest, new card version, snapshot 第 2 版."""
    d = clone(h1a_done)
    live = live_server(d)
    page.goto(live.u("/backtest"))
    page.locator("[data-job-start]").first.click()
    page.wait_for_url(re.compile(r"toast="), timeout=180_000)
    assert "BT-0002" in page.locator('[data-concl="history"]').inner_text()
    con = sqlite3.connect(str(d / "oppo.sqlite3"))
    try:
        seqs = [r[0] for r in con.execute("SELECT seq FROM snapshot WHERE month='2026-03' ORDER BY seq")]
    finally:
        con.close()
    assert seqs == [1, 2]
    _no_errors(page)
