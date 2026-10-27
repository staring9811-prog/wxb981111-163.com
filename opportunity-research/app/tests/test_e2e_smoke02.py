"""SMOKE-02: on an empty database, wen does the whole month in the browser — import H1-A, batch-confirm fields,
batch-confirm the "标题主干相同" merge suggestions, run compute and backtest, open the ranking and export.
Also RULE-06's last clause: the browser never requests anything outside 127.0.0.1."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from oppo.core import backup as BK
from tests.conftest import fixture_pkg, wait_job

SHOTS: Path | None = None  # set by the visual capture run to keep screenshots


def _job_done(live, page) -> None:  # type: ignore[no-untyped-def]
    st = wait_job(live, timeout=600)
    assert st.get("status") == "success", st
    page.wait_for_load_state("networkidle")


def month_in_browser(live, page) -> None:  # type: ignore[no-untyped-def]
    """SMOKE-02 steps 1–6 through the pages (shared with RULE-04)."""
    # 1 导入 H1-A
    page.goto(live.u("/packages?import=1"))
    page.fill("#pkg-path", str(fixture_pkg("H1-A")))
    page.click("[data-stage-path]")
    page.wait_for_selector("[data-import-submit]")
    page.click("[data-import-submit]")
    page.wait_for_url("**/packages/*/report*", timeout=120_000)
    assert "P-0001" in page.locator("h1").inner_text()

    # 2 按建议批量确认字段并保存为 v1
    page.goto(live.u("/fields"))
    page.check('[data-select-all="#field-table"]')
    page.click('button[data-open="#batch-modal"]')
    page.uncheck('#batch-form input[name="suggested_only"]')
    page.click('#batch-form button[type="submit"]')
    page.wait_for_load_state("networkidle")
    page.click('button[data-open="#save-modal"]')
    page.click('#save-modal button[type="submit"]')
    page.wait_for_load_state("networkidle")
    page.goto(live.u("/fields"))
    assert "v1" in page.locator(".topbar").inner_text()

    # 3 批量确认“标题主干相同”的合并建议
    page.goto(live.u("/spu/merge"))
    page.click('#panel-pending .chip:has-text("标题主干相同")')
    page.click('[data-select-all="#pending-table"]')  # selects the visible (filtered) rows; header goes indeterminate
    assert page.locator('#pending-table tbody tr:not([hidden]) input[data-row-check]:checked').count() >= 1
    page.click('#panel-pending button[data-open="#batch-modal"]')
    page.click('#batch-modal button[type="submit"]')
    page.wait_for_load_state("networkidle")

    # 4 运行计算
    page.goto(live.u("/signals"))
    page.click('button[data-job-start="/jobs/compute"]')
    _job_done(live, page)

    # 5 运行回测
    page.goto(live.u("/backtest"))
    page.locator('button[data-job-start^="/jobs/backtest"]').first.click()
    _job_done(live, page)
    page.goto(live.u("/card"))
    assert "已达标" in page.locator("main").inner_text()

    # 6 本月排序 → 导出
    page.goto(live.u("/ranking"))
    assert "第 1 版" in page.locator(".page-head").inner_text()
    page.click('.page-head a[data-open="#export-modal"]')
    page.click("[data-export-start]")
    page.wait_for_selector('[data-file="html"]', timeout=300_000)
    _job_done(live, page)


def test_smoke02_month_in_browser(live_server, data_dir, page, tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    import socket

    attempts: list[str] = []  # RULE-06: no outbound connection from the app during the whole month
    real_connect = socket.socket.connect

    def guarded(self, address):  # type: ignore[no-untyped-def]
        host = address[0] if isinstance(address, tuple) else str(address)
        if isinstance(address, tuple) and host not in ("127.0.0.1", "::1", "localhost"):
            attempts.append(str(address))
            raise OSError("outbound connections are not allowed in this test")
        return real_connect(self, address)

    monkeypatch.setattr(socket.socket, "connect", guarded)
    live = live_server(data_dir)
    with live.app.engine.begin() as conn:  # precondition: a writable backup directory
        ok, _ = BK.set_directory(conn, str(tmp_path / "备份"), "wen")
        assert ok
    outside: list[str] = []
    page.on("request", lambda r: outside.append(r.url) if not r.url.startswith(live.url) else None)

    month_in_browser(live, page)

    with live.app.engine.begin() as conn:
        exp = conn.execute(text("SELECT * FROM export_record ORDER BY id DESC")).mappings().first()
        bk = conn.execute(text("SELECT * FROM backup_record ORDER BY id DESC")).mappings().first()
        cv = conn.execute(text("SELECT MAX(version_no) FROM contract_version")).scalar()
        card = conn.execute(text("SELECT label, status FROM decision_card ORDER BY id DESC")).first()
        snap = conn.execute(text("SELECT month, seq FROM snapshot ORDER BY id DESC")).first()
    assert cv == 1 and tuple(card) == ("v1", "已达标") and tuple(snap) == ("2026-03", 1)
    assert Path(exp["xlsx_path"]).exists() and Path(exp["html_path"]).exists()
    assert len(load_workbook(exp["xlsx_path"], read_only=True).sheetnames) == 9
    assert bk["status"] == "success" and Path(bk["path"]).exists()

    # 7 工作台 8 步全部完成
    page.goto(live.u("/"))
    statuses = page.locator(".step").evaluate_all("els => els.map(e => e.dataset.status)")
    assert statuses == ["done"] * 8, statuses
    if SHOTS:
        page.screenshot(path=str(SHOTS / "smoke02_workbench.png"), full_page=True)
    assert not page.js_errors, page.js_errors
    assert not outside, outside
    assert not attempts, attempts
