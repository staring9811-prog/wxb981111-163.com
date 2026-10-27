"""Monthly export (F12: REQ-F12-01/02, AC-F12-01..03, AC-F05-08, AC-F05-11, AC-F07-11, DEC-006).

One export = one Excel workbook (9 sheets) + one self-contained offline HTML report, both built from the
same saved snapshot. Files are written under temporary names and renamed only when both are complete, so a
failure never leaves a half-written file. The automatic backup runs after the export commits
(see oppo.service.export).
"""

from __future__ import annotations

import hashlib
import os
import re
import uuid
from functools import lru_cache, partial
from pathlib import Path
from typing import Any, Callable

from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup
from sqlalchemy import select
from sqlalchemy.engine import Connection

from oppo import clock, fmt
from oppo import db as D
from oppo.core import backtest as BT
from oppo.core import cards as CARD
from oppo.core import categories as CAT
from oppo.core import compute as CMP
from oppo.core import ledger as LG
from oppo.core import packages as PK
from oppo.core import ranking as RK
from oppo.core import runs

SHEETS = ("首页", "细分排序", "锦囊", "判定卡版本", "回测摘要", "台账-快照", "台账-立项记录", "台账-对账", "版本信息")
SHEET_NOTES = {
    "细分排序": "本月全部细分：排名、得分、预期命中率区间、样本量与锦囊动作（顺序与本月排序页一致）",
    "锦囊": "当前判定卡下的规则：建议动作、触发条件、预期命中率区间、依据与状态",
    "判定卡版本": "每个判定卡版本的入选信号、阈值、权重与回测依据",
    "回测摘要": "生成当前判定卡的回测：切分、前 20% 命中率、倍数、对照、干扰因素与局限",
    "台账-快照": "全部月度快照、导出与对账状态",
    "台账-立项记录": "立项决定、补录与更正记录",
    "台账-对账": "到期快照的竞品新品对账与规则对账",
    "版本信息": "本次导出所用的数据包、字段约定、品类配置、判定卡、回测与快照版本",
}
WEB = Path(__file__).resolve().parents[1] / "web"


class ExportError(Exception):
    pass


Report = Callable[[str, float], None]


# ---- shared content -------------------------------------------------------------------------------


def status_line(card: dict[str, Any] | None) -> tuple[str, str]:
    """(kind, sentence) for the first line of both files (AC-F12-02, AC-F07-11)."""
    if card is None:
        return "warn", CARD.UNVERIFIED_TEXT
    if card["status"] == "未验证":
        b = card.get("basis") or {}
        reason = "：".join(x for x in (b.get("reason_class"), b.get("reason_text")) if x)
        return "warn", CARD.UNVERIFIED_TEXT + (f"（未达标：{reason}）" if reason else "")
    if card["status"] == "达标（截面）":
        return "cross", CARD.CROSS_TEXT
    return "ok", f"判定卡 {card['label']} 已达标（历史口径）"


def _recon_texts(conn: Connection, cid: int, cfg: dict[str, Any], current_id: int, now: str) -> list[dict[str, Any]]:
    window = int(cfg["success"]["window_months"])
    recs = {r["snapshot_id"]: r for r in LG.reconciliations(conn, cid)}
    out = []
    latest_T = {m: (CMP.latest_run(conn, cid, m) or {}).get("t_month") for m in ("history", "cross")}
    for s in RK.snapshot_rows(conn, cid):
        exported = s["exported_at"]
        if s["id"] == current_id:
            exported = now
        due = PK.add_months(s["month"], window)
        used = RK.reconcile_version(conn, cid, s["month"], s["mode"])
        if s["id"] in recs:
            r = recs[s["id"]]
            text = f"已对账（{r['created_at'][:10]}）" + ("（截面口径）" if s["mode"] == "cross" else "")
            state = "已对账"
        elif used and used["id"] != s["id"] and s["id"] != current_id:
            text = "不用于对账" + ("（同月已有导出版）" if used.get("exported_at") else "（同月已有更新的版本）")
            state = "不用于对账"
        elif latest_T.get(s["mode"]) and latest_T[s["mode"]] >= due:
            text, state = "可对账", "可对账"
        else:
            text, state = f"未到期（预计 {due}）", "未到期"
        out.append({**s, "exported_at": exported, "export_text": (f"已导出（{exported[:16]}）" if exported else "未导出"),
                    "recon_text": text, "recon_state": state, "due": due if state == "未到期" else None})
    return out


def _decisions(conn: Connection, cid: int) -> list[dict[str, Any]]:
    out = []
    for r in LG.records(conn, cid):
        snap = r.get("snapshot")
        row = None
        if snap:
            row = next((x for x in snap["content"].get("rows") or [] if x["segment_key"] == r.get("segment_key")), None)
        launch = r.get("launch") or {}
        if r.get("corrected_by"):
            state = f"已更正（见 {r['corrected_by']['code']}）"
            relation = f"已被 {r['corrected_by']['code']} 更正"
        else:
            relation = f"更正 {r['corrects']['code']}" if r.get("corrects") else ""
            if launch:
                state = f"已上架 {launch.get('launch_month')}；自家对账：{launch.get('status')}"
            elif r.get("decision") == "不立项":
                state = "理由：" + (r.get("reason") or "无数据")
            else:
                state = "在研"
        kind = r.get("kind") or "记录"
        out.append({
            "code": r["code"], "project": r.get("project_name"), "decision": r.get("decision"),
            "segment": r.get("segment_name") or (r.get("segment_key") or "").replace("|", " · "),
            "reason": r.get("reason"), "snapshot": snap["label"] if snap else None,
            "snapshot_none": "补录，无当时快照" if kind == "补录" else "无对应快照",
            "source": kind + (f"（{relation}）" if relation and kind == "更正" else ""), "relation": relation,
            "date": r.get("project_date"), "state": state, "launch_status": launch.get("status"),
            "launch_month": launch.get("launch_month"), "rank": row["rank"] if row else None,
            "tips": "；".join(f"{t['action']} {t['rule']}" for t in row["tips"]) if row else "",
            "operator": r.get("created_by"), "created_at": r.get("created_at"),
            "launch": launch, "kind": kind,
        })
    return out


def _recons(conn: Connection, cid: int) -> list[dict[str, Any]]:
    out = []
    for r in LG.reconciliations(conn, cid):
        snap = r.get("snapshot") or {}
        top_rate = r["top_hits"] / r["top_n"] if r["top_n"] else None
        all_rate = r["all_hits"] / r["all_n"] if r["all_n"] else None
        rules = "；".join(f"{x['code']} {x['action']} {fmt.mult(x['multiple'])} {x['result']}" for x in r["rules"])
        cross = r["mode"] == "cross"
        result = (r["within"] or "无数据") + (f"；{rules}" if rules else "") + (f" · {CARD.CROSS_LINE_TEXT}" if cross else "")
        out.append({
            "snapshot": snap.get("label", "无数据"), "batch": r["batch"], "batch_n": (r["all_n"] or 0) + (r["incomplete"] or 0),
            "k": len(((snap.get("content") or {}).get("top20") or {}).get("segments") or []),
            "top_hits": r["top_hits"], "top_n": r["top_n"], "top_rate": top_rate, "all_hits": r["all_hits"],
            "all_n": r["all_n"], "all_rate": all_rate, "multiple": r["multiple"],
            "expected": fmt.ci_text([r["exp_lo"], r["exp_hi"]]), "within": r["within"], "rules": rules,
            "result": result, "created_at": r["created_at"], "cross": cross,
        })
    return out


def _own(decisions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for d in decisions:
        la = d["launch"]
        if not la or d.get("relation", "").startswith("已被"):
            continue
        parts = []
        if la.get("best_bsr") is not None:
            parts.append(f"最好小类 BSR {la['best_bsr']:g}")
        if la.get("max_sales") is not None:
            parts.append(f"最高月销售额 {fmt.money(la['max_sales'])}")
        note = f"关联细分当时排名第 {d['rank']}，锦囊 {d['tips']}" if d["rank"] else "补录，无当时快照"
        out.append({"code": d["code"], "launch_month": la.get("launch_month"), "status": la.get("status") or "无数据",
                    "detail": "，".join(parts) or "无数据", "snapshot_text": d["snapshot"] or "补录，无当时快照",
                    "note": note})
    return out


def gather(conn: Connection, category_id: int, snap: dict[str, Any], operator: str, now: str) -> dict[str, Any]:
    from oppo.web.report import present as P

    cat = conn.execute(select(D.category).where(D.category.c.id == category_id)).mappings().first()
    cfg = CAT.current_version(conn, category_id)["config"]
    content = snap["content"]
    mode = snap["mode"]
    card = CARD.get(conn, snap["card_id"]) if snap.get("card_id") else None
    bt_row = BT.get(conn, snap["backtest_id"]) if snap.get("backtest_id") else None
    other = BT.latest(conn, category_id, "cross" if mode == "history" else "history")
    kits = P.kits(conn, category_id, mode, snap)
    decisions = _decisions(conn, category_id)
    vers = content.get("versions") or {}
    pkg = PK.by_label(conn, vers["package"]) if vers.get("package") else None
    cv = None
    if vers.get("contract"):
        cv = conn.execute(select(D.contract_version).where(
            D.contract_version.c.category_id == category_id,
            D.contract_version.c.version_no == int(str(vers["contract"]).lstrip("v")))).mappings().first()
    cfgv = conn.execute(select(D.category_version).where(
        D.category_version.c.category_id == category_id,
        D.category_version.c.version_no == int(str(vers.get("config", "v1")).lstrip("v")))).mappings().first()
    return {
        "cat": dict(cat) if cat else {"name": "品类", "site": "US", "node": "", "owner": operator}, "cfg": cfg,
        "snap": snap, "content": content, "card": card, "bt": P.backtest_view(conn, bt_row),
        "other_bt": P.backtest_view(conn, other) if other else None, "kits": kits,
        "card_versions": P.card_versions(conn, category_id),
        "ledger": {"snapshots": _recon_texts(conn, category_id, cfg, snap["id"], now), "decisions": decisions,
                   "recons": _recons(conn, category_id), "own": _own(decisions)},
        "vinfo": {"package": P.package_info(conn, pkg["id"]) if pkg else None, "contract": dict(cv) if cv else None,
                  "config": dict(cfgv) if cfgv else None},
        "status": status_line(card), "operator": operator, "generated": now[:16], "confounders": BT.CONFOUNDERS,
    }


# ---- HTML report -------------------------------------------------------------------------------------


@lru_cache(maxsize=1)
def _symbols() -> dict[str, str]:
    js = (WEB / "static" / "icons.js").read_text(encoding="utf-8")
    return {m.group(1): m.group(0) for m in re.finditer(r'<symbol id="ph-([a-z-]+)"[^>]*>.*?</symbol>', js)}


def _sprite(html: str) -> str:
    names = sorted(set(re.findall(r'href="#ph-([a-z-]+)"', html)))
    syms = _symbols()
    body = "".join(syms[n] for n in names if n in syms)
    return f'<svg style="display:none" aria-hidden="true">{body}</svg>'


@lru_cache(maxsize=1)
def _env() -> Environment:
    env = Environment(loader=FileSystemLoader(str(WEB / "report")), autoescape=select_autoescape(["html"]),
                      trim_blocks=True, lstrip_blocks=True)
    from oppo.web.report import present as P

    env.globals.update(fmt=fmt, ibar=partial(fmt.ibar, standalone=True), etag=fmt.etag, badge=fmt.badge,
                       icon=fmt.icon, segname=fmt.segname, pct=fmt.pct, pct0=fmt.pct0, num=fmt.num, money=fmt.money,
                       mult=fmt.mult, CARD=CARD, P=P)
    return env


def render_html(data: dict[str, Any], html_name: str) -> str:
    css = (WEB / "report" / "report.css").read_text(encoding="utf-8")
    page = _env().get_template("report.html").render(**{**data, "html_name": html_name}, css=Markup(css),
                                                    sprite=Markup("@@SPRITE@@"))
    return page.replace("@@SPRITE@@", _sprite(page), 1)


# ---- Excel -------------------------------------------------------------------------------------------


def _nd(v: Any) -> Any:
    return "无数据" if v is None or (isinstance(v, float) and v != v) else v


def build_xlsx(data: dict[str, Any], path: Path, report: Report) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    snap, card, bt, content = data["snap"], data["card"], data["bt"], data["content"]
    kind, sentence = data["status"]
    fill = {"warn": "F4EEE6", "cross": "EFECF6", "ok": "E8F1EC"}[kind]
    cat = data["cat"]
    ident = (f"{cat['name']} · 排序月份 {snap['month']} · 快照第 {snap['seq']} 版 · "
             f"判定卡 {card['label'] if card else '未生成'} {card['status'] if card else ''}")
    PCT, MULT = "0.0%", "0.00"
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name: str, headers: list[str], repeat_status: bool = False) -> Any:
        ws = wb.create_sheet(name)
        first = f"{sentence} · {ident}" if (repeat_status and kind != "ok") else ident
        ws.append([first])
        ws["A1"].font = Font(bold=True)
        if repeat_status and kind != "ok":
            ws["A1"].fill = PatternFill("solid", fgColor=fill)
        ws.append(headers)
        for c in ws[2]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A3"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = max(10, min(40, len(h) * 2 + 4))
        return ws

    def fmt_cells(ws: Any, row: int, cols: dict[int, str]) -> None:
        for c, f in cols.items():
            cell = ws.cell(row=row, column=c)
            if isinstance(cell.value, (int, float)):
                cell.number_format = f

    # 1 首页
    report("写入 Excel：第 1 / 9 张工作表：首页", 0.12)
    ws = wb.create_sheet("首页")
    ws.append([sentence])
    ws["A1"].font = Font(bold=True, color="8F5300" if kind == "warn" else ("5E45A0" if kind == "cross" else "1B7343"))
    ws["A1"].fill = PatternFill("solid", fgColor=fill)
    ws.append(["项目", "内容"])
    for c in ws[2]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A3"
    cfg = data["cfg"]
    avg = content.get("average") or {}
    rows = [
        ("品类", f"亚马逊{'美国' if cat.get('site') == 'US' else cat.get('site')} · {cat['name']}，节点 {cat.get('node')}"),
        ("排序月份", f"{snap['month']}（数据至 {content.get('data_to')}）"),
        ("快照", f"第 {snap['seq']} 版 · {snap['created_at'][:16]} 保存 · 内容指纹 {snap['sha256']}"),
        ("判定卡", f"{card['label']} · {card['status']} · {'截面' if snap['mode'] == 'cross' else '历史'}口径" if card else "无"),
        ("回测", (f"{bt['code']} · 倍数 {fmt.mult(bt['multiple'])}（90% 区间 {_mci(bt['mult_ci'])}）· 检验组样本量 "
                 f"{bt['all_n']}") if bt else "无"),
        ("参照人群", content.get("reference")),
        ("成功线", f"上架 {cfg['success']['window_months']} 个月内小类 BSR ≤ {cfg['success']['bsr_max']:g}，或 SPU 任一月销售额 ≥ "
                f"{fmt.money(cfg['success']['gmv_min'])}" + (f"（{CARD.CROSS_LINE_TEXT}）" if snap["mode"] == "cross" else "")),
        ("全类目平均命中率", f"{fmt.pct(avg.get('rate'))}（{avg.get('hits')} / {avg.get('n')}）" if avg.get("n") else "无数据"),
        ("生成时间", data["generated"]),
        ("操作人", data["operator"]),
    ]
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["本月结论 · 前 5 个有效细分"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["排名", "细分", "预期命中率", "区间下限", "区间上限", "锦囊动作"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in [x for x in content["rows"] if x["effective"]][:5]:
        ws.append([r["rank"], r["segment"], _nd(r["point"]), _nd(r["ci"][0]), _nd(r["ci"][1]),
                   "；".join(f"{t['action']} {t['rule']}" for t in r["tips"]) or "无适用锦囊"])
        fmt_cells(ws, ws.max_row, {3: PCT, 4: PCT, 5: PCT})
    ws.append([])
    ws.append(["工作表目录"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for name in SHEETS[1:]:
        ws.append([name, SHEET_NOTES[name]])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    for c in "CDEF":
        ws.column_dimensions[c].width = 14
    ws["A1"].alignment = Alignment(wrap_text=False)

    # 2 细分排序
    report("写入 Excel：第 2 / 9 张工作表：细分排序", 0.2)
    dims = [d.get("label", d["key"]) for d in cfg.get("segment_dims") or []] or ["宽度档", "形态", "价格带"]
    ws = sheet("细分排序", ["排名", "细分", *dims, "得分", "预期命中率", "区间下限", "区间上限", "样本量", "样本不足",
                        "锦囊动作", "规则编号", "其他适用锦囊", "口径"], repeat_status=True)
    for r in content["rows"]:
        parts = r["segment_key"].split("|")
        t0 = r["tips"][0] if r["tips"] else None
        eff = r["effective"]
        ws.append([r["rank"] if eff else None, r["segment"], *parts[:len(dims)],
                   round(r["score"], 4) if r["score"] is not None else "无数据", _nd(r["point"]),
                   _nd(r["ci"][0]) if eff else None, _nd(r["ci"][1]) if eff else None, r["sample"],
                   "否" if eff else "是", (t0["action"] if t0 else "无适用锦囊") if eff else None,
                   t0["rule"] if t0 and eff else None,
                   "；".join(f"{t['rule']} {t['action']}" + (f"（{t['attr_note']}）" if t.get("attr_note") else "")
                            for t in (r["tips"][1:] if eff else r["tips"])) or None,
                   "截面" if snap["mode"] == "cross" else "历史"])
        n = len(dims)
        fmt_cells(ws, ws.max_row, {3 + n: "0.00", 4 + n: PCT, 5 + n: PCT, 6 + n: PCT})

    # 3 锦囊
    report("写入 Excel：第 3 / 9 张工作表：锦囊", 0.28)
    ws = sheet("锦囊", ["规则编号", "版本", "建议动作", "触发条件", "预期命中率下限", "预期命中率上限", "依据批次", "样本量",
                      "状态", "状态依据", "来源", "理由", "适用细分", "修改人", "修改时间"], repeat_status=True)
    for k in data["kits"]:
        ws.append([k["code"], f"v{k['version']}", k["action"] + (f"（建议：{k['advice']}）" if k["manual"] and k["advice"] else ""),
                   k["cond"] + (f"（{k['attr_note']}）" if k["attr_note"] else ""),
                   _nd(k["ci"][0] if k["ci"] else None), _nd(k["ci"][1] if k["ci"] else None),
                   f"{k['batches']}（{k['backtest']}）" if k["batches"] else "无数据", _nd(k["n"]), k["status"], k["note"],
                   k["origin"], k["rule"].get("reason") if k["origin"] == "人工" else None,
                   "；".join(s["name"] for s in k["segments"]), k["updated_by"], (k["updated_at"] or "")[:16]])
        fmt_cells(ws, ws.max_row, {5: PCT, 6: PCT})

    # 4 判定卡版本
    report("写入 Excel：第 4 / 9 张工作表：判定卡版本", 0.36)
    ws = sheet("判定卡版本", ["版本", "是否当前", "生成时间", "口径", "状态", "来源回测", "信号", "阈值", "权重",
                         "定规则组批次范围", "检验组批次范围", "检验组样本量", "倍数", "倍数区间下限", "倍数区间上限",
                         "与上一版相比"], repeat_status=True)
    from oppo.web.report import present as P

    for v in data["card_versions"]:
        c = v["card"]
        b = c.get("basis") or {}
        sigs = c.get("signals") or [{}]
        for s in sigs:
            mci = v["mult_ci"] or [None, None]
            ws.append([v["label"], "是" if v["current"] else "否", c["created_at"][:16], "截面" if c["mode"] == "cross" else "历史",
                       c["status"] + ("（已被历史结论取代）" if v["superseded"] else ""), v["bt"],
                       s.get("label", "无入选信号"), P.threshold_text(s["signal"], s.get("threshold")) if s.get("signal") else None,
                       s.get("weight"), P.span(b.get("rule_batches")), P.span(b.get("test_batches")),
                       (f"新品-月 {v['model_rows']} 行" if c["mode"] == "cross" and v["model_rows"] else _nd(v["test_n"])),
                       _nd(v["multiple"]), _nd(mci[0]), _nd(mci[1]), v["diff"]])
            fmt_cells(ws, ws.max_row, {9: "0.00", 13: MULT, 14: MULT, 15: MULT})

    # 5 回测摘要
    report("写入 Excel：第 5 / 9 张工作表：回测摘要", 0.44)
    ws = sheet("回测摘要", ["项目", "内容", "", "", "", "", ""])
    for bv, title in ((bt, None), (data["other_bt"], "另一口径的结论（判定卡状态以历史为准）")):
        if not bv:
            continue
        if title:
            ws.append([])
            ws.append([title])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for item in _bt_summary(bv):
            ws.append(list(item))
        ws.append([])
        ws.append(["对照"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append(["排序依据", "前 20% 命中率", "分子", "分母", "区间下限", "区间上限", "说明"])
        for r in bv["rows"]:
            if r.get("group"):
                continue
            if r["ok"]:
                ws.append([r["label"], r["rate"], r["hits"], r["n"], _nd((r["ci"] or [None])[0]),
                           _nd((r["ci"] or [None, None])[1]), "倍数基准" if r["kind"] == "base" else
                           f"倍数 {fmt.mult(r.get('multiple'))}"])
                fmt_cells(ws, ws.max_row, {2: PCT, 5: PCT, 6: PCT})
            else:
                ws.append([r["label"], "无数据", "无数据", "无数据", "无数据", "无数据", f"{r['reason']}，不参与比较"])
        if bv["confounders"]:
            ws.append([])
            ws.append(["控制干扰因素前后（优势比）"])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.append(["信号", "控制前", "控制后"])
            for c in bv["confounders"]:
                ws.append([c["label"], _nd(c["before"]), _nd(c["after"])])
                fmt_cells(ws, ws.max_row, {2: MULT, 3: MULT})
        ws.append([])
        ws.append(["局限", bv["limitation"]])
        ws.append(["被拒绝的快照字段"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append(["字段名", "原因"])
        for f in bv["rejected"] or []:
            ws.append([f"{f['file_label']} · {f['field']}", "字段约定标为快照，回测不用" + (f"（{f['note']}）" if f.get("note") else "")])
        if not bv["rejected"]:
            ws.append(["无", ""])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 60

    # 6 台账-快照
    report("写入 Excel：第 6 / 9 张工作表：台账-快照", 0.52)
    ws = sheet("台账-快照", ["快照", "判定卡版本", "判定卡状态", "口径", "生成时间", "是否导出", "最后导出时间", "对账状态",
                          "预计可对账月份", "内容指纹"])
    for s in data["ledger"]["snapshots"]:
        ws.append([s["label"], s["card_label"], s["card_status"], "截面" if s["mode"] == "cross" else "历史", s["created_at"][:16],
                   "是" if s["exported_at"] else "否", (s["exported_at"] or "")[:16] or None, s["recon_state"], s["due"],
                   s["sha256"]])
    if not data["ledger"]["snapshots"]:
        ws.append(["暂无记录"])

    # 7 台账-立项记录
    report("写入 Excel：第 7 / 9 张工作表：台账-立项记录", 0.6)
    ws = sheet("台账-立项记录", ["编号", "立项名称", "决定", "关联细分", "理由", "对应快照", "来源", "更正关系", "立项日期",
                            "上架状态", "上架月", "自家对账结果", "当时排名", "当时锦囊", "操作人", "记录时间"])
    for d in data["ledger"]["decisions"]:
        ws.append([d["code"], d["project"], d["decision"], d["segment"], d["reason"], d["snapshot"], d["kind"], d["relation"] or None,
                   d["date"], "已上架" if d["launch_month"] else ("不适用" if d["decision"] == "不立项" else "在研"),
                   d["launch_month"], d["launch_status"], d["rank"], d["tips"] or None, d["operator"],
                   (d["created_at"] or "")[:16]])
    if not data["ledger"]["decisions"]:
        ws.append(["暂无记录"])

    # 8 台账-对账
    report("写入 Excel：第 8 / 9 张工作表：台账-对账", 0.68)
    ws = sheet("台账-对账", ["快照", "对应批次", "起量新 SPU", "前 20% 细分数", "实际命中", "实际命中分母", "实际命中率",
                          "全类目命中", "全类目分母", "全类目命中率", "倍数", "回测预期区间", "是否落在区间内", "规则对账",
                          "对账时间", "口径"])
    for r in data["ledger"]["recons"]:
        ws.append([r["snapshot"], r["batch"], r["batch_n"], r["k"], r["top_hits"], r["top_n"], _nd(r["top_rate"]),
                   r["all_hits"], r["all_n"], _nd(r["all_rate"]), _nd(r["multiple"]), r["expected"],
                   {"落在区间内": "是"}.get(r["within"] or "", "否" if r["within"] else "无数据"), r["rules"] or None,
                   (r["created_at"] or "")[:16], CARD.CROSS_LINE_TEXT if r["cross"] else "历史"])
        fmt_cells(ws, ws.max_row, {7: PCT, 10: PCT, 11: MULT})
    if not data["ledger"]["recons"]:
        ws.append(["暂无记录"])

    # 9 版本信息
    report("写入 Excel：第 9 / 9 张工作表：版本信息", 0.74)
    ws = sheet("版本信息", ["项目", "版本", "时间", "指纹或说明"])
    vi = data["vinfo"]
    pkg = vi["package"]
    vers = content.get("versions") or {}
    ws.append(["数据包", pkg["label"] if pkg else vers.get("package"),
               (pkg.get("created_at") or "")[:16] if pkg else None,
               f"{'截面' if pkg and pkg.get('kind') == 'cross' else '历史'} · 覆盖 {pkg['months']} · 指纹 {pkg.get('fingerprint')}"
               if pkg else None])
    ws.append(["字段约定", vers.get("contract"), (vi["contract"] or {}).get("created_at", "")[:16] or None, None])
    ws.append(["品类配置", vers.get("config"), (vi["config"] or {}).get("created_at", "")[:16] or None,
               f"负责人 {cat.get('owner')}"])
    ws.append(["判定卡", card["label"] if card else None, card["created_at"][:16] if card else None,
               card["status"] if card else None])
    ws.append(["回测", bt["code"] if bt else None, bt["started"] if bt else None, f"随机种子 {bt['seed']}" if bt else None])
    ws.append(["快照", f"{snap['month']} 第 {snap['seq']} 版", snap["created_at"][:16], snap["sha256"]])
    ws.append(["导出时间", data["generated"], None, None])
    ws.append(["操作人", data["operator"], None, None])
    ws.append(["文件", data.get("xlsx_name"), None, data.get("html_name")])
    wb.save(str(path))


def _mci(ci: Any) -> str:
    if not ci or ci[0] is None or ci[1] is None:
        return "无数据"
    return f"{ci[0]:.2f}–{ci[1]:.2f}"


def _bt_summary(bv: dict[str, Any]) -> list[tuple[str, Any]]:
    from oppo.web.report import present as P

    out: list[tuple[str, Any]] = [
        ("回测编号", bv["code"]), ("口径", "截面" if bv["cross"] else "历史"),
        ("开始时间与用时", f"{bv['started']} · {bv['duration']}"), ("随机种子", bv["seed"]),
        ("所用版本", f"数据包 {bv['package']['label'] if bv['package'] else '无数据'} · 字段约定 v{bv['contract_no']} · "
                 f"品类配置 v{bv['config_no']}"),
        ("已判定批次", f"{len(bv['decided'])} 个（{P.span(bv['decided'])}）"),
    ]
    if bv["test_batches"]:
        rule_n = (f"新品-月 {bv['model_rows'].get('model')} 行" if bv["cross"] else f"新 SPU {bv['rule_counts'].get('spus')} 个")
        test_n = (f"新品-月 {bv['model_rows'].get('test')} 行 · 新 SPU {bv['test_counts'].get('spus')} 个" if bv["cross"]
                  else f"新 SPU {bv['test_counts'].get('spus')} 个")
        out += [("建模组" if bv["cross"] else "定规则组", f"{P.span(bv['rule_batches'])} · {rule_n}"),
                ("检验组", f"{P.span(bv['test_batches'])} · {test_n}"), ("前 20% 细分数", bv["k"])]
    if bv["rate"] is not None:
        out += [("前 20% 命中率", f"{bv['top_hits']} / {bv['top_n']} = {fmt.pct(bv['rate'])}（Wilson 90% {fmt.ci_text(bv['top_ci'], 1)}）"),
                ("全类目平均", f"{bv['all_hits']} / {bv['all_n']} = {fmt.pct(bv['avg'])}"),
                ("倍数", f"{fmt.mult(bv['multiple'])}（重抽样 90% 区间 {_mci(bv['mult_ci'])}）")]
    out += [("结论", bv["conclusion"]), ("未达标原因与数值", "：".join(x for x in (bv["reason_class"], bv["reason_text"]) if x)
                                       or "不适用"), ("样本不足细分数", len(bv["insufficient"]))]
    return out


# ---- the job -----------------------------------------------------------------------------------------


def _free_names(folder: Path, base: str) -> tuple[Path, Path]:
    n = 1
    while True:
        stem = base if n == 1 else f"{base}_{n}"
        x, h = folder / f"{stem}.xlsx", folder / f"{stem}.html"
        if not x.exists() and not h.exists():
            return x, h
        n += 1


def export_month(conn: Connection, settings: Any, category_id: int, operator: str, month: str | None, mode: str,
                 report: Report) -> dict[str, Any]:
    started = clock.now_iso()
    report("读取快照", 0.03)
    snaps = RK.snapshot_rows(conn, category_id, mode)
    if month:
        snaps = [s for s in snaps if s["month"] == month]
    if not snaps:
        raise ExportError("还没有可导出的排序快照")
    snap = snaps[0]
    if hashlib.sha256(snap["content_json"].encode("utf-8")).hexdigest() != snap["sha256"]:
        raise ExportError(f"快照 {snap['label']} 的内容指纹与保存时不一致，已停止导出")
    now = clock.now_iso()
    data = gather(conn, category_id, snap, operator, now)
    folder = Path(settings.exports_dir) / snap["month"]
    try:
        folder.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportError(f"导出目录无法创建：{folder}（{exc.strerror or exc}）") from exc
    base = f"{data['cat']['name']}_机会判定_{snap['month']}_快照{snap['seq']}"
    xlsx, html = _free_names(folder, base)
    data["xlsx_name"], data["html_name"] = xlsx.name, html.name
    tag = uuid.uuid4().hex[:8]
    tmp_x, tmp_h = folder / f".{xlsx.name}.{tag}.tmp", folder / f".{html.name}.{tag}.tmp"
    done: list[Path] = []
    try:
        build_xlsx(data, tmp_x, report)
        report("生成网页报告", 0.8)
        tmp_h.write_text(render_html(data, html.name), encoding="utf-8")
        report("写入文件", 0.9)
        os.replace(tmp_x, xlsx)
        done.append(xlsx)
        os.replace(tmp_h, html)
        done.append(html)
        eid = conn.execute(D.export_record.insert().values(
            category_id=category_id, month=snap["month"], snapshot_id=snap["id"], xlsx_path=str(xlsx),
            html_path=str(html), xlsx_size=xlsx.stat().st_size, html_size=html.stat().st_size,
            created_by=operator, created_at=now)).inserted_primary_key[0]
        conn.execute(D.snapshot_export.insert().values(snapshot_id=snap["id"], export_id=eid, created_at=now))
        runs.record_step(conn, category_id, "导出", started, clock.now_iso(), operator,
                         result=f"{snap['label']} · {xlsx.name}", link="/ranking", month=snap["month"])
    except PermissionError as exc:
        _cleanup([tmp_x, tmp_h, *done])
        raise ExportError(f"{Path(exc.filename or xlsx).name} 无法写入（可能正被其他程序打开），关闭后重试") from exc
    except Exception:
        _cleanup([tmp_x, tmp_h, *done])
        raise
    report("自动备份", 0.95)
    return {"export_id": int(eid), "xlsx_path": str(xlsx), "html_path": str(html), "xlsx_size": xlsx.stat().st_size,
            "html_size": html.stat().st_size, "snapshot_id": snap["id"], "month": snap["month"],
            "snapshot": snap["label"], "card_status": snap["card_status"]}


def _cleanup(paths: list[Path]) -> None:
    for p in paths:
        try:
            if p.exists():
                p.unlink()
        except OSError:
            pass

