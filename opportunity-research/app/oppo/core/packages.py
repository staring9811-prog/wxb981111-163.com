"""Import and validate data packages (F01: REQ-F01-01..03, DEC-018..021)."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import stat
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from sqlalchemy import func, select, text
from sqlalchemy.engine import Connection

from oppo import clock
from oppo import db as D
from oppo.core import spec as S

MANIFEST_NAMES = ("清单.json", "manifest.json")
KIND_ALIASES = {"历史": "history", "history": "history", "截面": "cross", "cross": "cross"}
MONTH_RE = re.compile(r"^(\d{4})(\d{2})$")
DATE_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


class ImportRejected(Exception):
    """Raised for duplicate imports (AC-F01-05); carries the existing label."""

    def __init__(self, label: str):
        super().__init__(f"该数据包已导入：{label}")
        self.label = label


@dataclass
class ImportResult:
    package_id: int
    label: str
    conclusion: str
    kind: str | None
    duplicate_of: str | None = None
    messages: list[str] = field(default_factory=list)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def month_of(value: str) -> str | None:
    m = MONTH_RE.match(value.strip())
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    if not (1 <= mo <= 12 and 1990 <= y <= 2100):
        return None
    return f"{y:04d}-{mo:02d}"


def date_of(value: str) -> str | None:
    m = DATE_RE.match(value.strip())
    if not m:
        return None
    try:
        pd.Timestamp(year=int(m.group(1)), month=int(m.group(2)), day=int(m.group(3)))
    except ValueError:
        return None
    return value.strip()


def num_of(value: str) -> float | None:
    try:
        v = float(value.strip())
    except ValueError:
        return None
    if v != v:  # NaN literal
        return None
    return v


def bool_of(value: str) -> int | None:
    v = value.strip().lower()
    if v in ("是", "1", "true", "y", "yes"):
        return 1
    if v in ("否", "0", "false", "n", "no"):
        return 0
    return None


PARSERS = {S.MONTH: month_of, S.DATE: date_of, S.NUM: num_of, S.BOOL: bool_of}


def infer_dtype(values: pd.Series) -> str:
    vals = [v for v in values if isinstance(v, str) and v.strip() != ""]
    if not vals:
        return S.TEXT
    sample = vals[:2000]
    if all(num_of(v) is not None for v in sample):
        return S.NUM
    if all(date_of(v) is not None for v in sample):
        return S.DATE
    if all(month_of(v) is not None for v in sample):
        return S.MONTH
    return S.TEXT


def add_months(month: str, n: int) -> str:
    y, m = int(month[:4]), int(month[5:7])
    idx = y * 12 + (m - 1) + n
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def month_range(a: str, b: str) -> list[str]:
    out, cur = [], a
    while cur <= b:
        out.append(cur)
        cur = add_months(cur, 1)
    return out


def _read_table(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        df = pd.read_excel(path, dtype=str, keep_default_na=False)
    else:
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    return df.fillna("")


def _next_label(conn: Connection) -> str:
    n = conn.execute(select(func.count()).select_from(D.package)).scalar() or 0
    return f"P-{n + 1:04d}"


@dataclass
class FileCheck:
    code: str
    spec: S.FileSpec | None
    path: Path | None
    df: pd.DataFrame | None = None
    rows: int = 0
    items: list[dict[str, Any]] = field(default_factory=list)
    excluded: set[int] = field(default_factory=set)  # 0-based data row indexes
    parsed: dict[str, list[Any]] = field(default_factory=dict)
    extra_fields: list[str] = field(default_factory=list)
    optional_present: list[str] = field(default_factory=list)


def _row_label(i: int) -> int:
    """File line number of data row i (header is line 1)."""
    return i + 2


def _grade(bad: int, total: int) -> str:
    if bad == 0:
        return "pass"
    return "warn" if total and bad / total <= 0.01 else "block"


def _check_file(fc: FileCheck) -> None:
    spec, df = fc.spec, fc.df
    assert spec is not None and df is not None
    total = len(df)
    fc.rows = total
    cols = set(df.columns)
    missing_cols = [f for f in spec.required if f not in cols]
    for f in spec.required:
        if f in missing_cols:
            fc.items.append(dict(file_code=fc.code, field=f, check_name="required_field", result="block",
                                 affected_rows=total, affected_note="整列", message=f"缺少必需字段“{f}”"))
        else:
            fc.items.append(dict(file_code=fc.code, field=f, check_name="required_field", result="pass",
                                 affected_rows=0, message=None))
    fc.optional_present = [f for f in spec.optional if f in cols]
    known = set(spec.required) | set(spec.optional)
    fc.extra_fields = [c for c in df.columns if c not in known]
    if missing_cols:
        return
    typed = {f: t for f, t in {**spec.required, **spec.optional}.items() if f in cols}
    bad_rows: set[int] = set()
    bad_by_field: dict[str, list[int]] = {}
    for f, t in typed.items():
        series = df[f]
        parser = PARSERS.get(t)
        if parser is None:
            fc.parsed[f] = [v.strip() if v.strip() != "" else None for v in series]
            continue
        out: list[Any] = []
        for i, raw in enumerate(series):
            if raw.strip() == "":
                out.append(None)
                continue
            v = parser(raw)
            if v is None:
                bad_rows.add(i)
                bad_by_field.setdefault(f, []).append(i)
            out.append(v)
        fc.parsed[f] = out
    grade = _grade(len(bad_rows), total)
    sample = [_row_label(i) for i in sorted(bad_rows)[:5]]
    fields_txt = "、".join(bad_by_field) if bad_by_field else None
    # Per-field breakdown for the report (rows and up to 3 example raw values per field).
    detail = None
    if bad_by_field:
        detail = D.dumps({"by_field": {f: len(ix) for f, ix in bad_by_field.items()},
                          "examples": {f: list(dict.fromkeys(df[f].iloc[ix[:20]].str.strip()))[:3]
                                       for f, ix in bad_by_field.items()}})
    fc.items.append(dict(
        file_code=fc.code, field=fields_txt, check_name="parse", result=grade, affected_rows=len(bad_rows),
        affected_note=detail, sample_rows_json=D.dumps(sample),
        excluded_rows_json=D.dumps([_row_label(i) for i in sorted(bad_rows)]) if grade == "warn" else None,
        message=(None if not bad_rows else
                 f"{len(bad_rows)} 行无法解析（{fields_txt}），占 {len(bad_rows) / total:.1%}" +
                 ("，计算中排除这些行" if grade == "warn" else "，超过 1%，不能计算")),
    ))
    # Primary key uniqueness: later occurrences of a key count like unparsable rows (DEC-019).
    key_cols = list(spec.key)
    keys = df[key_cols].apply(lambda r: "\x1f".join(v.strip() for v in r), axis=1)
    dup_mask = keys.duplicated(keep="first")
    dup_rows = set(int(i) for i in dup_mask[dup_mask].index)
    dgrade = _grade(len(dup_rows), total)
    fc.items.append(dict(
        file_code=fc.code, field="×".join(key_cols), check_name="unique_key", result=dgrade,
        affected_rows=len(dup_rows), sample_rows_json=D.dumps([_row_label(i) for i in sorted(dup_rows)[:5]]),
        excluded_rows_json=D.dumps([_row_label(i) for i in sorted(dup_rows)]) if dgrade == "warn" else None,
        message=None if not dup_rows else f"主键重复 {len(dup_rows)} 行（保留首次出现的行）",
    ))
    fc.excluded = bad_rows | dup_rows


def _coverage(files: dict[str, FileCheck], kind: str, manifest: dict[str, Any]) -> None:
    cov = manifest.get("覆盖月份") or manifest.get("观察月份") or {}
    m_from, m_to = cov.get("起"), cov.get("止")
    if not (m_from and m_to):
        return
    def months_of(code: str) -> set[str]:
        fc = files.get(code)
        if not fc or fc.df is None or "月份" not in fc.parsed:
            return set()
        return {m for i, m in enumerate(fc.parsed["月份"]) if m and i not in fc.excluded}
    c1 = "C1" if kind == "history" else "S1"
    c2 = "C2" if kind == "history" else "S2"
    c3 = "C3" if kind == "history" else "S3"
    need = month_range(m_from, m_to)
    have = months_of(c1)
    missing = [m for m in need if m not in have]
    _cov_item(files, c1, missing, f"{m_from} 至 {m_to}",
              "缺少月份：" + "、".join(missing) + "；涉及月份的信号为“无数据”" if missing else None)
    fc2, fc3 = files.get(c2), files.get(c3)
    if fc2 and fc3 and fc2.df is not None and fc3.df is not None and "上架日期" in fc2.parsed and "月份" in fc3.parsed:
        launch = {c: d for c, d in zip(fc2.parsed.get("子ASIN", []), fc2.parsed["上架日期"]) if c and d}
        have3: dict[str, set[str]] = {}
        for i, (c, m) in enumerate(zip(fc3.parsed.get("子ASIN", []), fc3.parsed["月份"])):
            if c and m and i not in fc3.excluded:
                have3.setdefault(c, set()).add(m)
        short = 0
        months_missing = 0
        for c, d in launch.items():
            start = max(d[:7], m_from)
            end = min(add_months(d[:7], 5), m_to)
            gap = [m for m in month_range(start, end) if m not in have3.get(c, set())]
            if gap:
                short += 1
                months_missing += len(gap)
        _cov_item(files, c3, [1] * short, "每个新品上架起至少 6 个月或至最新月",
                  f"{short} 个新品子体缺 {months_missing} 个月的轨迹；缺月份的新 SPU 可能判为“数据不全”" if short else None)
        lo = min(launch.values())[:7] if launch else None
        _cov_item(files, c2, [], f"上架日 {lo} 起" if lo else "无新品", None)
    fc4 = files.get("C4")
    if fc4 and fc4.df is not None and "子ASIN" in fc4.parsed:
        tagged = set(c for c in fc4.parsed["子ASIN"] if c)
        children = set()
        for code in (c1, c2):
            fc = files.get(code)
            if fc and fc.df is not None:
                children |= set(c for c in fc.parsed.get("子ASIN", []) if c)
        untagged = sorted(children - tagged)
        _cov_item(files, "C4", untagged, f"覆盖 {c1}、{c2} 全部子 ASIN",
                  f"{len(untagged)} 个子 ASIN 没有属性标签，归入各维度“未知”档" if untagged else None)
    fc5 = files.get("C5")
    if fc5 and fc5.df is not None and "月份" in fc5.parsed:
        have5 = months_of("C5")
        lo5 = min(have5) if have5 else None
        _cov_item(files, "C5", [], f"{lo5} 至 {max(have5)}" if have5 else "无数据", None)
    fc6 = files.get("C6")
    if fc6 and fc6.df is not None and "评论日期" in fc6.parsed:
        dates = [d for i, d in enumerate(fc6.parsed["评论日期"]) if d and i not in fc6.excluded]
        _cov_item(files, "C6", [], f"至 {max(dates)[:7]}" if dates else "无评论", None)
    if "C7" in files and files["C7"].df is not None:
        _cov_item(files, "C7", [], "不适用", None, na=True)


def _cov_item(files: dict[str, FileCheck], code: str, missing: list[Any], note: str, msg: str | None,
              na: bool = False) -> None:
    fc = files.get(code)
    if fc is None:
        return
    fc.items.append(dict(file_code=code, field=None, check_name="coverage",
                         result="na" if na else ("warn" if missing else "pass"),
                         affected_rows=len(missing) if missing else 0, affected_note=note, message=msg))


def _field_evidence(fc: FileCheck) -> list[dict[str, Any]]:
    df = fc.df
    assert df is not None and fc.spec is not None
    out = []
    typed = {**fc.spec.required, **fc.spec.optional}
    child_col = "子ASIN" if "子ASIN" in df.columns else None
    has_month = "月份" in df.columns
    for col in df.columns:
        series = df[col]
        nonempty = series.str.strip() != ""
        dtype = typed.get(col) or infer_dtype(series)
        ev: dict[str, Any] = {"nonnull_share": float(nonempty.mean()) if len(series) else 0.0,
                              "rows": int(len(series))}
        ex = next((v for v in series if v.strip()), None)
        ev["example"] = ex[:60] if ex else None
        if child_col and has_month and col not in (child_col, "月份"):
            g = df[nonempty][[child_col, col]].groupby(child_col)[col].nunique()
            if len(g):
                ev["same_all_months_share"] = float((g <= 1).mean())
                ev["same_all_months_n"] = int(len(g))
        out.append({"field": col, "dtype": dtype, "evidence": ev})
    return out


def _make_readonly(path: Path) -> None:
    os.chmod(path, stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)


def _materialize(src: Path, workdir: Path) -> Path:
    if src.is_dir():
        return src
    if src.suffix.lower() == ".zip":
        with zipfile.ZipFile(src) as z:
            z.extractall(workdir)
        entries = [p for p in workdir.iterdir()]
        # A zip that wraps one top-level folder.
        if len(entries) == 1 and entries[0].is_dir() and not any((workdir / m).exists() for m in MANIFEST_NAMES):
            return entries[0]
        return workdir
    raise ValueError("请选择 zip 文件或文件夹")


def import_package(conn: Connection, settings: Any, category_id: int, src: Path, operator: str) -> ImportResult:
    """Run the full SEQ-F01-01: fingerprint → duplicate check → validate → read-only copy → parse."""
    from oppo.core import categories as CAT
    from oppo.core import contract, health, runs, spu

    started = clock.now_iso()
    cfg_ver = CAT.current_version(conn, category_id)
    cfg = cfg_ver["config"]
    tmp_root = settings.data_dir / "tmp"
    tmp_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=tmp_root) as tmp:
        root = _materialize(Path(src), Path(tmp))
        manifest_path = next((root / m for m in MANIFEST_NAMES if (root / m).exists()), None)
        all_files = sorted(p for p in root.iterdir() if p.is_file())
        hashes = {p.name: sha256_file(p) for p in all_files}
        fingerprint = hashlib.sha256("\n".join(f"{h}" for h in sorted(hashes.values())).encode()).hexdigest()
        # AC-F01-05: identical package (every file fingerprint equal) → no new version.
        dup = conn.execute(select(D.package.c.label).where(
            D.package.c.category_id == category_id, D.package.c.fingerprint == fingerprint)).first()
        if dup:
            raise ImportRejected(dup[0])

        label = _next_label(conn)
        items: list[dict[str, Any]] = []
        manifest: dict[str, Any] = {}
        kind: str | None = None
        files: dict[str, FileCheck] = {}
        if manifest_path is None:
            items.append(dict(file_code=None, field="清单文件", check_name="manifest", result="block",
                              affected_note="整个数据包", message="清单文件缺失"))
        else:
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
                kind = KIND_ALIASES.get(str(manifest.get("类型", "")).strip())
            except (json.JSONDecodeError, UnicodeDecodeError):
                manifest, kind = {}, None
            if kind is None:
                items.append(dict(file_code=None, field="类型", check_name="manifest", result="block",
                                  affected_note="整个数据包", message="类型无法识别（应为“历史”或“截面”）"))
            elif (manifest.get("品类") and manifest.get("品类") != cfg.get("name")) or \
                    (manifest.get("站点") and manifest.get("站点") != cfg.get("site")):
                items.append(dict(file_code=None, field="品类/站点", check_name="manifest", result="block",
                                  affected_note="整个数据包",
                                  message=f"清单中的品类或站点（{manifest.get('品类')} · {manifest.get('站点')}）"
                                          f"与当前品类配置（{cfg.get('name')} · {cfg.get('site')}）不一致"))
                kind = None
            else:
                items.append(dict(file_code=None, field="清单文件", check_name="manifest", result="pass",
                                  message=f"类型：{'历史数据包（约定 v1）' if kind == 'history' else '截面数据包（约定 v1.1）'}"))
        if kind is not None:
            dims = tuple(cfg.get("attribute_dims") or S.C4_DIMS)
            specs = S.history_specs(dims) if kind == "history" else S.cross_specs(dims)
            declared = manifest.get("文件") or {}
            for fs in specs:
                name = declared.get(fs.code)
                path = root / name if name else None
                if path is None or not path.exists():
                    guess = [p for p in all_files if p.name.upper().startswith(fs.code)]
                    path = guess[0] if guess else None
                fc = FileCheck(fs.code, fs, path)
                files[fs.code] = fc
                if path is None:
                    fc.items.append(dict(file_code=fs.code, field=None, check_name="required_file",
                                         result="block" if fs.required_file else "na",
                                         affected_note="整个文件",
                                         message="缺少必需文件" if fs.required_file else "未提供（可选）"))
                    continue
                fc.items.append(dict(file_code=fs.code, field=None, check_name="required_file", result="pass",
                                     message=path.name))
                try:
                    fc.df = _read_table(path)
                except Exception as exc:  # unreadable file is a block
                    fc.items.append(dict(file_code=fs.code, field=None, check_name="parse", result="block",
                                         affected_note="整个文件", message=f"文件无法读取：{exc}"))
                    fc.df = None
                    continue
                _check_file(fc)
            if not any(i["result"] == "block" for f in files.values() for i in f.items):
                _coverage(files, kind, manifest)
            for fc in files.values():
                items.extend(fc.items)

        conclusion = "block" if any(i["result"] == "block" for i in items) else (
            "warn" if any(i["result"] == "warn" for i in items) else "pass")
        cov = manifest.get("覆盖月份") or manifest.get("观察月份") or {}
        now = clock.now_iso()
        pid = conn.execute(D.package.insert().values(
            category_id=category_id, label=label, kind=kind, manifest_json=D.dumps(manifest) if manifest else None,
            source_name=Path(src).name, months_from=cov.get("起"), months_to=cov.get("止"), conclusion=conclusion,
            fingerprint=fingerprint, config_version_id=cfg_ver["id"], created_by=operator, created_at=now,
        )).inserted_primary_key[0]
        for p in all_files:
            code = next((c for c, fc in files.items() if fc.path == p), "清单" if p == manifest_path else "其他")
            conn.execute(D.package_file.insert().values(
                package_id=pid, file_code=code, filename=p.name, sha256=hashes[p.name],
                row_count=files[code].rows if code in files else None,
                optional=int(code in files and files[code].spec is not None and not files[code].spec.required_file)))
        for it in items:
            conn.execute(D.validation_item.insert().values(package_id=pid, **it))

        result = ImportResult(pid, label, conclusion, kind)
        if conclusion != "block":
            dest = settings.packages_dir / label
            dest.mkdir(parents=True, exist_ok=True)
            for p in all_files:
                target = dest / p.name
                shutil.copy2(p, target)
                _make_readonly(target)
            conn.execute(D.package.update().where(D.package.c.id == pid).values(readonly_path=str(dest)))
            for fc in files.values():
                if fc.df is None:
                    continue
                for ev in _field_evidence(fc):
                    conn.execute(D.package_field.insert().values(
                        package_id=pid, file_code=fc.code, field=ev["field"], dtype=ev["dtype"],
                        evidence_json=D.dumps(ev["evidence"])))
                _store_raw(conn, pid, fc, manifest)
            key = "current_history_package" if kind == "history" else "current_cross_package"
            D.set_meta(conn, f"{key}:{category_id}", str(pid))
            contract.on_import(conn, category_id, pid, operator)
            spu.refresh_suggestions(conn, category_id, pid, operator)
        runs.record_step(conn, category_id, "导入校验", started, clock.now_iso(), operator,
                         result={"pass": "通过", "warn": "有警告", "block": "阻断"}[conclusion] + f" · {label}",
                         link=f"/packages/{pid}/report", package_id=pid)
        if conclusion != "block":
            hstart = clock.now_iso()
            health.run(conn, category_id, pid)
            runs.record_step(conn, category_id, "数据体检", hstart, clock.now_iso(), operator,
                             result="完成", link=f"/health?package={pid}", package_id=pid)
        return result


def _store_raw(conn: Connection, pid: int, fc: FileCheck, manifest: dict[str, Any]) -> None:
    df = fc.df
    assert df is not None and fc.spec is not None
    keep = [i for i in range(len(df)) if i not in fc.excluded]
    if not keep:
        return
    P = fc.parsed

    def col(name: str) -> list[Any]:
        vals = P.get(name)
        if vals is None:
            return [None] * len(keep)
        return [vals[i] for i in keep]

    extra = fc.extra_fields
    extra_json = [D.dumps({e: df.at[i, e] for e in extra if df.at[i, e] != ""}) if extra else None for i in keep]
    base = {"package_id": pid, "row_no": [_row_label(i) for i in keep], "extra_json": extra_json}
    code = fc.code
    if code in ("C1", "S1"):
        out = pd.DataFrame({
            **base, "file_code": code, "month": col("月份"), "child": col("子ASIN"), "parent": col("当月父ASIN"),
            "brand": col("品牌"), "seller": col("卖家"), "title": col("标题"), "launch_date": col("上架日期"),
            "child_units": col("子体月销量"), "child_sales": col("子体月销售额"), "parent_units": col("父体月销量"),
            "bsr": col("小类BSR"), "node": col("类目节点"), "fulfillment": col("配送方式"),
            "variant_count": col("变体数"), "first_review": col("首次评分日期"), "spu_no": col("SPU编号"),
            "bought_label": col("过去一个月购买"), "price": col("当月价格"), "rating": col("当月评分"),
            "bullets": col("五点描述"),
        })
        out.to_sql("raw_c1", conn, if_exists="append", index=False, chunksize=20000)
    elif code in ("C2", "S2"):
        out = pd.DataFrame({
            **base, "file_code": code, "child": col("子ASIN"), "parent": col("父ASIN"), "brand": col("品牌"),
            "title": col("标题"), "launch_date": col("上架日期"), "fulfillment": col("配送方式"),
            "variant_count": col("变体数"), "first_review": col("首次评分日期"), "spu_no": col("SPU编号"),
            "bullets": col("五点描述"),
        })
        out.to_sql("raw_c2", conn, if_exists="append", index=False, chunksize=20000)
    elif code in ("C3", "S3"):
        out = pd.DataFrame({
            **base, "file_code": code, "month": col("月份"), "child": col("子ASIN"),
            "child_units": col("子体月销量"), "child_sales": col("子体月销售额"), "bsr": col("小类BSR"),
            "parent_units": col("父体月销量"),
        })
        out.to_sql("raw_c3", conn, if_exists="append", index=False, chunksize=20000)
    elif code == "C4":
        dims = [c for c in fc.spec.required if c != "子ASIN" and not c.endswith("证据状态")]
        rows = []
        children = col("子ASIN")
        for j, i in enumerate(keep):
            for d in dims:
                rows.append({"package_id": pid, "row_no": _row_label(i), "extra_json": None, "child": children[j],
                             "dim": d, "value": P[d][i], "status": P[f"{d}证据状态"][i]})
        pd.DataFrame(rows).to_sql("raw_c4", conn, if_exists="append", index=False, chunksize=20000)
    elif code == "C5":
        out = pd.DataFrame({**base, "keyword": col("关键词"), "month": col("月份"), "volume": col("月搜索量"),
                            "tags": col("关键词属性标签"), "purchases": col("月购买量")})
        out.to_sql("raw_c5", conn, if_exists="append", index=False, chunksize=20000)
    elif code == "C6":
        out = pd.DataFrame({**base, "review_id": col("评论ID"), "asin": col("ASIN"), "level": col("归属级别"),
                            "review_date": col("评论日期"), "stars": col("星级"), "pain_code": col("痛点事件代码"),
                            "structural": col("是否结构类痛点"), "label_version": col("标注版本")})
        out.to_sql("raw_c6", conn, if_exists="append", index=False, chunksize=20000)
    elif code == "C7":
        out = pd.DataFrame({**base, "name": col("立项名称"), "project_date": col("立项日期"),
                            "segment_or_attr": col("关联细分或属性"), "launch_date": col("上架日期"),
                            "asin": col("对应ASIN")})
        out.to_sql("raw_c7", conn, if_exists="append", index=False, chunksize=20000)


# ---- preview before import (PAGE-F01-01 清单预览) ------------------------------------------------


def _count_rows(path: Path) -> int | None:
    """Data rows (header excluded) without parsing values."""
    try:
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            n = max(0, (wb.active.max_row or 1) - 1)
            wb.close()
            return n
        import csv

        with open(path, newline="", encoding="utf-8-sig") as f:
            return max(0, sum(1 for _ in csv.reader(f)) - 1)
    except Exception:  # noqa: BLE001 - an unreadable file shows as 无数据 in the preview
        return None


def preview(conn: Connection, settings: Any, category_id: int, src: Path) -> dict[str, Any]:
    """Read the manifest and list the files of a package without importing it (AC-F01-05/06)."""
    from oppo.core import categories as CAT

    cfg = CAT.current_version(conn, category_id)["config"]
    src = Path(src)
    out: dict[str, Any] = {"source_name": src.name, "manifest_found": False, "kind": None, "kind_raw": None,
                           "error": None, "files": [], "extra_files": [], "duplicate": None,
                           "config": {"name": cfg.get("name"), "site": cfg.get("site"), "node": cfg.get("node")}}
    if src.is_file():
        out["size"] = src.stat().st_size
    elif src.is_dir():
        out["size"] = sum(p.stat().st_size for p in src.iterdir() if p.is_file())
    tmp_root = settings.data_dir / "tmp"
    tmp_root.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory(dir=tmp_root) as tmp:
            root = _materialize(src, Path(tmp))
            manifest_path = next((root / m for m in MANIFEST_NAMES if (root / m).exists()), None)
            all_files = sorted(p for p in root.iterdir() if p.is_file())
            hashes = {p.name: sha256_file(p) for p in all_files}
            fingerprint = hashlib.sha256("\n".join(sorted(hashes.values())).encode()).hexdigest()
            out["fingerprint"] = fingerprint
            dup = conn.execute(select(D.package.c.id, D.package.c.label, D.package.c.created_at).where(
                D.package.c.category_id == category_id, D.package.c.fingerprint == fingerprint)).first()
            if dup:
                out["duplicate"] = {"id": dup[0], "label": dup[1], "created_at": dup[2]}
            manifest: dict[str, Any] = {}
            if manifest_path is None:
                out["error"] = "missing"
            else:
                out["manifest_found"] = True
                try:
                    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
                except (json.JSONDecodeError, UnicodeDecodeError):
                    manifest = {}
                out["kind_raw"] = str(manifest.get("类型", "")).strip()
                out["kind"] = KIND_ALIASES.get(out["kind_raw"])
                if out["kind"] is None:
                    out["error"] = "kind"
            out["manifest"] = manifest
            cov = manifest.get("覆盖月份") or manifest.get("观察月份") or {}
            out["months_from"], out["months_to"] = cov.get("起"), cov.get("止")
            if out["months_from"] and out["months_to"]:
                out["months_n"] = len(month_range(out["months_from"], out["months_to"]))
            out["category"], out["site"], out["node"] = manifest.get("品类"), manifest.get("站点"), manifest.get("类目节点")
            if out["kind"] and ((out["category"] and out["category"] != cfg.get("name")) or
                                (out["site"] and out["site"] != cfg.get("site"))):
                out["error"] = "category"
            used: set[str] = set()
            if out["kind"]:
                dims = tuple(cfg.get("attribute_dims") or S.C4_DIMS)
                specs = S.history_specs(dims) if out["kind"] == "history" else S.cross_specs(dims)
                declared = manifest.get("文件") or {}
                for fs in specs:
                    name = declared.get(fs.code)
                    path = root / name if name else None
                    if path is None or not path.exists():
                        guess = [p for p in all_files if p.name.upper().startswith(fs.code)]
                        path = guess[0] if guess else None
                    row = {"code": fs.code, "label": S.FILE_NAMES.get(fs.code, fs.code), "required": fs.required_file,
                           "declared": name, "present": path is not None, "filename": path.name if path else None,
                           "rows": _count_rows(path) if path else None,
                           "sha": hashes.get(path.name) if path else None}
                    if path is not None:
                        used.add(path.name)
                    out["files"].append(row)
            out["extra_files"] = [p.name for p in all_files if p.name not in used and p != manifest_path]
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        out["error"] = "unreadable"
        out["unreadable"] = f"{src.name} 无法读取：{getattr(exc, 'strerror', None) or exc}"
    req = [f for f in out["files"] if f["required"]]
    out["required_total"] = len(req)
    out["required_present"] = sum(1 for f in req if f["present"])
    out["optional_present"] = [f["code"] for f in out["files"] if not f["required"] and f["present"]]
    out["missing_required"] = [f["code"] for f in req if not f["present"]]
    out["can_import"] = out["error"] is None and out["duplicate"] is None
    return out


def file_codes_text(codes: list[str]) -> str:
    """C1–C7 / S1–S3、C4 style range text for the package list."""
    order = ["C1", "C2", "C3", "C4", "C5", "C6", "C7"]
    if all(c.startswith("C") for c in codes):
        idx = sorted(order.index(c) for c in codes if c in order)
        if idx and idx == list(range(idx[0], idx[-1] + 1)):
            return f"{order[idx[0]]}–{order[idx[-1]]}" if len(idx) > 1 else order[idx[0]]
        return "、".join(order[i] for i in idx)
    s = [c for c in codes if c.startswith("S")]
    rest = [c for c in codes if not c.startswith("S")]
    head = f"{s[0]}–{s[-1]}" if len(s) > 1 else "".join(s)
    return "、".join([head] + sorted(rest, key=lambda c: order.index(c) if c in order else 9))


# ---- queries -------------------------------------------------------------------------------


def current_package(conn: Connection, category_id: int, kind: str) -> dict[str, Any] | None:
    key = "current_history_package" if kind == "history" else "current_cross_package"
    v = D.get_meta(conn, f"{key}:{category_id}")
    if not v:
        return None
    row = conn.execute(select(D.package).where(D.package.c.id == int(v))).mappings().first()
    return dict(row) if row else None


def get(conn: Connection, package_id: int) -> dict[str, Any] | None:
    row = conn.execute(select(D.package).where(D.package.c.id == package_id)).mappings().first()
    return dict(row) if row else None


def by_label(conn: Connection, label: str) -> dict[str, Any] | None:
    row = conn.execute(select(D.package).where(D.package.c.label == label)).mappings().first()
    return dict(row) if row else None


def list_packages(conn: Connection, category_id: int) -> list[dict[str, Any]]:
    rows = conn.execute(select(D.package).where(D.package.c.category_id == category_id)
                        .order_by(D.package.c.id.desc())).mappings().all()
    return [dict(r) for r in rows]


def files(conn: Connection, package_id: int) -> list[dict[str, Any]]:
    return [dict(r) for r in conn.execute(select(D.package_file).where(D.package_file.c.package_id == package_id)
                                          .order_by(D.package_file.c.id)).mappings().all()]


def validation(conn: Connection, package_id: int) -> list[dict[str, Any]]:
    rows = conn.execute(select(D.validation_item).where(D.validation_item.c.package_id == package_id)
                        .order_by(D.validation_item.c.id)).mappings().all()
    out = []
    for r in rows:
        d = dict(r)
        d["sample_rows"] = D.loads(r["sample_rows_json"], [])
        d["excluded_rows"] = D.loads(r["excluded_rows_json"], [])
        out.append(d)
    return out


def raw_row_counts(conn: Connection, package_id: int) -> dict[str, int]:
    out: dict[str, int] = {}
    for tbl, col in (("raw_c1", "file_code"), ("raw_c2", "file_code"), ("raw_c3", "file_code")):
        for code, n in conn.execute(text(f"SELECT {col}, COUNT(*) FROM {tbl} WHERE package_id=:p GROUP BY {col}"),
                                    {"p": package_id}):
            out[code] = int(n)
    for tbl, code in (("raw_c5", "C5"), ("raw_c6", "C6"), ("raw_c7", "C7")):
        n = conn.execute(text(f"SELECT COUNT(*) FROM {tbl} WHERE package_id=:p"), {"p": package_id}).scalar()
        if n:
            out[code] = int(n)
    n4 = conn.execute(text("SELECT COUNT(DISTINCT child) FROM raw_c4 WHERE package_id=:p"), {"p": package_id}).scalar()
    if n4:
        out["C4"] = int(n4)
    return out


def assert_computable(conn: Connection, package_id: int) -> None:
    """AC-F01-02: a package with a block can't start any computation."""
    pkg = get(conn, package_id)
    if pkg is None:
        raise ValueError("数据包不存在")
    if pkg["conclusion"] == "block":
        raise PermissionError(f"数据包 {pkg['label']} 存在阻断项，不能计算")
