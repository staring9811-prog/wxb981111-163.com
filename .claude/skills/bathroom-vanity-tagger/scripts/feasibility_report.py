"""Answer "which fields can the current data tag, which cannot yet".

Reads tag_cached.py output (tagging_summary/evidence/work_queues.json) and the
project payload, writes a feasibility workbook and a Markdown summary.
Network-free; counts only, no new tagging.

  python3 -I feasibility_report.py --tag-dir OUT --data-root DATA \
      --out-xlsx 可行性.xlsx --out-md 可行性.md [--base-evidence v2/tagging_evidence.json]
"""
import argparse,json,sys
from collections import Counter
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))  # works under python -I
from tag_cached import FEASIBILITY,HEADERS
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

FIELDS=HEADERS[:14]
# What each field can be read from, and what unblocks the rest.
FIELD_NOTES={
    'D1产品形态':('标题/SKU/详情Mounting Type/五点的floating、wall-mounted、freestanding','未写形态的行看主图即可判断悬浮/落地'),
    'D2安装方式':('五点/详情Assembly字段的assembled、assembly required','缺子体详情的补采asin_detail；父体候选多为同系列口径，补采后可升级'),
    'D3盆数':('SKU Single/Dual、标题single/double sink','多数标题只写with Sink不写数量；看主图可判断单/双盆'),
    'D4台面形态':('vessel、undermount、integrated、drop-in等原文','看主图可区分台上/台下/一体盆；integrated单词不证明同材一体'),
    'D5元素':('标题/五点中的Fluted、Shaker、Barn Door等风格词','元素本质是外观，文本只覆盖少数；需逐子体看图'),
    'D6柜体颜色':('SKU Color、详情Color、标题末尾颜色','组合色（如Black White）需分清柜/盆/五金；冲突与未写的行看图'),
    'D7柜体材质':('详情Material、五点的MDF/plywood/solid wood等','内部板材图片无法证明；父体候选需子体详情确认，未写的需厂家规格'),
    'D8台面材质':('详情Material/Top Material、五点中绑定top的材料词','近半缺子体详情需补采；仿石纹理图片不能证明材质；未写的需厂家规格'),
    'D8台盆材质':('详情Material/五点中绑定sink/basin的材料词','缺子体详情的补采；未写的需厂家规格'),
    'D9宽(in)':('SKU Size、标题宽度、详情Product Dimensions','缺子体详情的补采；冲突多为卖家标签互换或名义/实际尺寸差'),
    'D9深(in)':('详情Product Dimensions（带D/W/H标签）；商品尺寸为无方向三维数','确认DIM-DWH-ORDER/DIM-DW-SWAP规则后可补大部分；其余补采'),
    'D9高(in)':('详情Product Dimensions（带D/W/H标签）；商品尺寸为无方向三维数','确认DIM-DWH-ORDER规则后可补大部分；SKU“3H”是3孔龙头位不是高度'),
    'D10镜子':('五点/标题中mirror included / not included等明确表述；未提及按用户规则默认无','默认无是业务口径，不是文本证据；按子体口径分析时用feasibility分开A/A2'),
    'D11龙头':('五点/标题中faucet included / not included、标题列出Faucet等；未提及按用户规则默认无','pre-drilled faucet hole不代表含龙头；默认无是业务口径，用feasibility分开A/A2'),
}

def verdict(r,n):
    tagged=r['A']+r['A2']+r.get('A3',0)
    if tagged>=0.75*n:return '可打标'
    if tagged+r['D']>=0.75*n:return '确认规则后可打标'
    if tagged>=0.35*n:return '部分可打标'
    return '现阶段不行（大部分）'

def pct(x,n):return '%.1f%%'%(100*x/n) if n else '-'

def style(ws,widths,wrap_cols=()):
    head=PatternFill('solid',fgColor='DDE7F3')
    for c in ws[1]:c.font=Font(bold=True);c.fill=head;c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    for col in wrap_cols:
        for c in ws[get_column_letter(col)][1:]:c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--tag-dir',required=True);ap.add_argument('--data-root',required=True)
    ap.add_argument('--out-xlsx',required=True);ap.add_argument('--out-md',required=True)
    ap.add_argument('--base-evidence',help='上一版tagging_evidence.json，用于列出变更')
    ap.add_argument('--dim-rules-dir',help='用--enable-dim-rules跑出的目录，用于给出启用后的覆盖')
    args=ap.parse_args()
    tag=Path(args.tag_dir)
    summary=json.loads((tag/'tagging_summary.json').read_text(encoding='utf-8'))
    rows=json.loads((tag/'tagging_evidence.json').read_text(encoding='utf-8'))['rows']
    queues=json.loads((tag/'tagging_work_queues.json').read_text(encoding='utf-8'))
    payload=json.loads((Path(args.data_root)/'sales_payload.json').read_text(encoding='utf-8'))['detailRows']
    dim=json.loads((Path(args.dim_rules_dir)/'tagging_summary.json').read_text(encoding='utf-8')) if args.dim_rules_dir else None
    target=[(e,r) for e,r in zip(rows,payload) if e['includeInBathroomAnalysis']]
    n=len(target);asins=len({e['childASIN'] for e,r in target})
    wb=Workbook()

    ws=wb.active;ws.title='字段结论'
    ws.append(['字段','现阶段结论','可打标行(子体证据+业务规则+父体补充)','占比','其中父体补充','仅父体候选','有冲突','确认规则后可补','缺子体详情需补采','需看图','子体文本未写','可打标唯一子ASIN','对照目录可打标占比','主要数据源','卡点与下一步'])
    lines=[]
    for h in FIELDS:
        r=summary['feasibility'][h]['rows'];u=summary['feasibility'][h]['uniqueASINs']
        after=''
        if dim:
            d=dim['feasibility'][h]['rows'];after=pct(d['A']+d['A2']+d.get('A3',0),n)
        v=verdict(r,n);src,nxt=FIELD_NOTES[h]
        tg=r['A']+r['A2']+r.get('A3',0)
        ws.append([h,v,tg,pct(tg,n),r.get('A3',0),r['B'],r['C'],r['D'],r['E'],r['F'],r['G'],u['A']+u['A2']+u.get('A3',0),after,src,nxt])
        lines.append((h,v,r,after,nxt))
    style(ws,[12,16,14,8,10,10,8,12,14,9,11,12,14,46,52],(14,15))

    ws=wb.create_sheet('数据源覆盖')
    cov=summary['textSourceCoverage']
    ws.append(['数据源','粒度','覆盖行','覆盖唯一子ASIN','能支撑的字段','局限'])
    ws.append(['SKU属性',  '子体',sum(bool(r.get('SKU属性')) for e,r in target),len({e['childASIN'] for e,r in target if r.get('SKU属性')}),'颜色、宽度、盆数(Single/Dual)','变体名可能是营销名；“3H”等编码需识别'])
    ws.append(['商品标题','子体',n,asins,'形态、宽度、盆型、龙头/镜子是否列出','标题常是SEO关键词堆叠，with Sink不代表盆数'])
    ws.append(['子体详情（overviews+五点）','子体',cov['childDetail']['rows'],cov['childDetail']['uniqueASINs'],'材质、颜色、带标签尺寸、安装、包含/排除','只有907个子ASIN有精确子体详情；D/W标签可能被卖家标反'])
    ws.append(['卖家精灵导出“详细参数/亮点”','子体（导出月快照）',cov['sellerSpriteExport']['rows'],cov['sellerSpriteExport']['uniqueASINs'],'Material、Color、Product Dimensions等','只覆盖导出的代表ASIN；v3新增使用'])
    ws.append(['父体详情五点填充','父体',cov['parentFeatures']['rows'],cov['parentFeatures']['uniqueASINs'],'安装、材质、包含/排除（仅父体候选）','不能证明某个子体的尺寸/材质/配置'])
    ws.append(['父体详情overviews（材质等）','父体','—','—','材质（仅父体候选）','代表变体口径；颜色/尺寸不采用'])
    ws.append(['商品尺寸（无方向a×b×c）','子体',sum(bool(r.get('商品尺寸')) for e,r in target),len({e['childASIN'] for e,r in target if r.get('商品尺寸')}),'深、高（需确认顺序规则）','无D/W/H标签；按D×W×H顺序在交叉核对样本上一致率约99%'])
    ws.append(['主图链接','子体',sum(bool(r.get('主图链接')) for e,r in target),len({e['childASIN'] for e,r in target if r.get('主图链接')}),'形态、盆数、盆型、元素、颜色（需看图）','链接可访问但本轮未看图；图片不能证明内部板材或台面真材'])
    ws.append(['流量词/评论','子体','—','—','不作属性证据','词数不是属性；评论未采集'])
    ws.append(['两类子体文本都没有','—',cov['noChildDetailOrExport']['rows'],cov['noChildDetailOrExport']['uniqueASINs'],'只剩SKU+标题+父体候选','需补采asin_detail（约每个子ASIN一次调用）'])
    style(ws,[30,16,10,14,36,52],(5,6))

    ws=wb.create_sheet('逐行可行性')
    ws.append(['站点','目标月份','子ASIN','核验父ASIN','商品标题','有子体详情','有卖家精灵导出','五点来自父体']+FIELDS)
    for e,r in target:
        src=e['textSources']
        ws.append([e['marketplace'],e['month'],e['childASIN'],e['verifiedParentASIN'],(r.get('商品标题') or '')[:120],'是' if src['childDetail'] else '否','是' if src['sellerSpriteExport'] else '否','是' if src['parentFeatures'] else '否']+[e['feasibility'][h] for h in FIELDS])
    style(ws,[6,9,13,13,48,9,10,9]+[8]*14)

    ws=wb.create_sheet('尺寸规则候选')
    ws.append(['目标月份','子ASIN','已选宽(in)','候选深(in)','候选高(in)','规则','原文（商品尺寸/详情）'])
    for e,r in target:
        fd=e['fields'].get('D9深(in)',{});fh=e['fields'].get('D9高(in)',{})
        cd=fd.get('swapCorrectedCandidate') or fd.get('orderInferenceCandidate');ch=fh.get('orderInferenceCandidate')
        if not (cd or ch):continue
        rules='、'.join(sorted({x['ruleID'] for x in (cd,ch) if x}))
        ws.append([e['month'],e['childASIN'],e['tags']['D9宽(in)'],cd['label'] if cd else None,ch['label'] if ch else None,rules,(cd or ch)['quote'][:120]])
    style(ws,[9,13,10,10,10,26,60])

    ws=wb.create_sheet('补采子体详情清单')
    ws.append(['子ASIN','核验父ASIN','目标月份','被卡字段(缺子体详情)'])
    rec=set(queues['recollectChildDetailASINs']);info={}
    for e,r in target:
        if e['childASIN'] not in rec:continue
        x=info.setdefault(e['childASIN'],{'parent':set(),'months':set(),'fields':set()})
        x['parent'].add(e['verifiedParentASIN']);x['months'].add(e['month']);x['fields']|={h for h in FIELDS if e['feasibility'][h]=='E'}
    for a in sorted(info):
        x=info[a];ws.append([a,'、'.join(sorted(x['parent'])),'、'.join(sorted(x['months'])),'、'.join(h for h in FIELDS if h in x['fields'])])
    style(ws,[13,26,30,70])

    ws=wb.create_sheet('看图核验队列')
    ws.append(['子ASIN','主图链接','商品链接','目标月份','需看图字段'])
    for q in queues['imageReviewQueue']:ws.append([q['childASIN'],q['mainImage'],q['productURL'],'、'.join(q['months']),'、'.join(q['fields'])])
    style(ws,[13,60,40,30,50])

    changes=[]
    if args.base_evidence:
        old=json.loads(Path(args.base_evidence).read_text(encoding='utf-8'))['rows']
        ws=wb.create_sheet('与上一版变更')
        ws.append(['字段','变更行数','原值','新值','行数'])
        for h in FIELDS+['打标状态']:
            c=Counter((str(x['tags'][h]),str(y['tags'][h])) for x,y in zip(old,rows) if y['includeInBathroomAnalysis'] and x['tags'][h]!=y['tags'][h])
            tot=sum(c.values());changes.append((h,tot,c.most_common(3)))
            for (a,b),k in c.most_common(8):ws.append([h,tot,a,b,k])
        style(ws,[14,10,34,34,8])

    ws=wb.create_sheet('图例与口径')
    ws.append(['代码','含义'])
    for k,v in FEASIBILITY.items():ws.append([k,v])
    for t in ['A/A2/A3计入“可打标”；A3是子体缺失时按用户2026-10-27规则用父体信息补充；A2是用户已确认的业务归类或推导（如marble归岗石、石英石台面配盆推陶瓷盆），不是接口直接返回。',
              'D5/D1/D3/D4/D6这类外观字段，文本没写时归“需看图”；材质、尺寸、包含关系文本没写时，有子体详情归“子体文本未写”，没有子体详情归“需补采”。',
              '范围：只统计includeInBathroomAnalysis=true的%d行/%d个唯一子ASIN；另有目标月归属待核验%d行、不纳入%d行，按口径不打标，不是数据问题。'%(n,asins,summary['rowStates'].get('目标月归属待核验',0),summary['rowStates'].get('不纳入浴室柜分析',0)),
              '可打标≠已验证准确：没有独立真值集，不报告准确率；属性是当前采集快照，不代表各月月末属性。',
              '图片：本轮独立看图0张。']:
        ws.append(['说明',t])
    style(ws,[8,120],(2,))
    wb.save(args.out_xlsx)

    md=['# 浴室柜打标可行性（%s）'%summary['ruleVersion'],'',
        '范围：%d 行子ASIN×月份，%d 个唯一子ASIN。规则版本 %s；新增接口调用 0；独立看图 0。'%(n,asins,summary['ruleVersion']),'',
        '| 字段 | 现阶段结论 | 可打标 | 其中父体补充 | 冲突 | 确认规则可补 | 需补采 | 需看图 | 文本未写 |'+(' 对照目录 |' if dim else ''),
        '|---|---|---|---|---|---|---|---|---|'+('---|' if dim else '')]
    for h,v,r,after,nxt in lines:
        tg=r['A']+r['A2']+r.get('A3',0)
        md.append('| %s | %s | %s (%s) | %s | %s | %s | %s | %s | %s |%s'%(h,v,tg,pct(tg,n),r.get('A3',0),r['C'],r['D'],r['E'],r['F'],r['G'],(' %s |'%after) if dim else ''))
    md+=['','## 卡点与下一步','']+['- **%s**：%s'%(h,nxt) for h,v,r,after,nxt in lines]
    md+=['','## 工作队列','','- 需补采子体详情（asin_detail）：%d 个子ASIN'%len(queues['recollectChildDetailASINs']),'- 需看图核验：%d 个子ASIN（主图链接可访问，未执行）'%len(queues['imageReviewQueue'])]
    if changes:
        md+=['','## 与上一版的变更（行数）','']+['- %s：%d 行；%s'%(h,t,'；'.join('%s→%s %d'%(a,b,k) for (a,b),k in top)) for h,t,top in changes if t]
    Path(args.out_md).write_text('\n'.join(md)+'\n',encoding='utf-8')
    print(json.dumps({'rows':n,'uniqueASINs':asins,'verdicts':{h:v for h,v,r,a,x in lines}},ensure_ascii=False))

if __name__=='__main__':main()
